"""Functional test for the Prompt Library spreadsheet: search + category filter vs. a Python reference.

Run: python3 Build_Tools/tests/test_prompt_library.py
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from build_prompt_library import LIB, STEM, load  # noqa: E402

RECALC = next(Path("/root/.claude/skills").rglob("xlsx/scripts/recalc.py"), None)
SRC = LIB / f"{STEM}_Searchable.xlsx"
CASES = [("invoice", "All"), ("scope", "All"), ("", "Tough Conversations"), ("proposal", "Proposals & Pricing"),
         ("zzzz-no-match", "All"), ("LEAD-0", "All"), ("testimonial", "Offboarding, Testimonials & Referrals")]


def expected(items, kw, cat):
    out = []
    for i in items:
        hay = " ".join([i["id"], i["category"], i["title"], i["objective"], i["use_case"], i["prompt"].strip()]).lower()
        if (cat == "All" or cat == i["category"]) and (kw == "" or kw.lower() in hay):
            out.append(i["id"])
    return out


def main():
    items = load()
    fails = 0
    tmp = Path(tempfile.mkdtemp())
    for n, (kw, cat) in enumerate(CASES):
        f = tmp / f"case{n}.xlsx"
        shutil.copy(SRC, f)
        wb = load_workbook(f)
        wb["Search"]["C4"] = kw or None
        wb["Search"]["F4"] = cat
        wb.save(f)
        res = json.loads(subprocess.run([sys.executable, str(RECALC), str(f), "90"], capture_output=True, text=True).stdout)
        assert res.get("status") == "success", res
        ws = load_workbook(f, data_only=True)["Search"]
        got = [ws[f"B{r}"].value for r in range(8, 108) if ws[f"B{r}"].value]
        exp = expected(items, kw, cat)
        ok = got == exp and ws["B5"].value == f"{len(exp)} matching prompt(s)"
        fails += not ok
        print(("PASS" if ok else "FAIL"), repr(kw), cat, len(got), "results" + ("" if ok else f" expected {len(exp)}: {exp[:5]} got {got[:5]} / {ws['B5'].value}"))
    cats = load_workbook(tmp / "case0.xlsx", data_only=True)["Categories"]
    total = cats["C3"].value
    per = sum(cats[f"C{r}"].value for r in range(4, 15))
    ok = total == 100 and per == 100
    fails += not ok
    print("PASS" if ok else "FAIL", "category counts", total, per)
    print(f"{len(CASES) + 1 - fails}/{len(CASES) + 1} passed")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
