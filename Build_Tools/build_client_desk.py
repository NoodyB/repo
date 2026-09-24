"""Build Product B: The Client Desk — a solo business command center workbook.

Works in Microsoft Excel (2010+), Google Sheets (upload & open) and LibreOffice Calc.
Only Excel-2007-era functions are used (IF, INDEX, MATCH, COUNTIFS, SUMIFS, RANK, CODE...).

Run: python3 Build_Tools/build_client_desk.py [--today YYYY-MM-DD] [--out path]
"""
import argparse
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from brand import ROOT

OUT_DIR = ROOT / "Digital_Products_Business" / "04_Premium_Template"
FILE_NAME = "BandOfOne_The-Client-Desk_v1.0.xlsx"

# ---------------------------------------------------------------- styles
INK, TEAL, DEEP, MIST, RUST, BUTTER, STONE, SLATE = "1B2430", "0F6B64", "0A4A45", "E3F1EE", "B8471F", "F6C85F", "E4DED3", "56606B"
INPUT_FILL = PatternFill("solid", fgColor="FFF8E1")   # yellow = you type here
CALC_FILL = PatternFill("solid", fgColor="F1F4F3")    # grey = calculated
HEAD_FILL = PatternFill("solid", fgColor=DEEP)
HELP_FILL = PatternFill("solid", fgColor="EDEBE6")
TILE_FILL = PatternFill("solid", fgColor=MIST)
WHITE = PatternFill("solid", fgColor="FFFFFF")
F = "Arial"
thin = Side(style="thin", color="D9D4CA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=F, bold=True, size=18, color=DEEP)
SUB_FONT = Font(name=F, italic=True, size=10, color=SLATE)
BODY = Font(name=F, size=10, color=INK)
BOLD = Font(name=F, size=10, color=INK, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATE_FMT = "yyyy-mm-dd"
MONEY_FMT = "#,##0.00;[Red]-#,##0.00;-"
TODAY = "Settings!$C$6"      # effective 'today' (override or TODAY())
TERMS = "Settings!$C$7"      # default payment terms (days)

STAGES = [("1 · Inquiry", 0.10, "Pipeline"), ("2 · Discovery", 0.25, "Pipeline"), ("3 · Proposal sent", 0.50, "Pipeline"),
          ("4 · Onboarding", 1.0, "Active"), ("5 · Delivery", 1.0, "Active"), ("6 · Wrap-up & payment", 1.0, "Active"),
          ("7 · Past client", 0.0, "Closed"), ("Lost", 0.0, "Closed")]
DECISIONS = ["Eliminate", "Template", "AI draft", "Automate", "Keep human"]
CATEGORIES = ["Getting clients", "Starting projects", "During projects", "Money", "Finishing", "Running the business"]
PROJ_STATUS = ["Not started", "In progress", "Waiting on client", "Complete", "On hold"]
SOURCES = ["Referral", "Website", "LinkedIn", "Instagram", "TikTok / YouTube", "Email list", "Marketplace", "Past client", "Other"]
SERVICES = ["Website design", "Website care plan", "Strategy session", "Other"]
REMINDERS = [("A", -3, "Heads-up (before due)"), ("B", 0, "Due today"), ("C", 7, "Gentle nudge"), ("D", 14, "Firm follow-up"), ("E", 30, "Final notice")]
WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

CLIENT_ROWS = (5, 154)
INVOICE_ROWS = (5, 204)
PROJECT_ROWS = (5, 104)
AUDIT_ROWS = (5, 44)
PROMPT_ROWS = (5, 124)
SOP_ROWS = (5, 54)
REVIEW_ROWS = (14, 65)


def style_range(ws, ref, font=None, fill=None, fmt=None, align=None, border=True):
    for row in ws[ref]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if fmt: c.number_format = fmt
            if align: c.alignment = align
            if border: c.border = BORDER


def header(ws, row, headers, widths, kinds):
    """kinds: 'in' input, 'calc' formula, 'help' helper column"""
    for i, (h, w, k) in enumerate(zip(headers, widths, kinds), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = H_FONT
        c.fill = HEAD_FILL if k != "help" else PatternFill("solid", fgColor=SLATE)
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 34


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    ws["A2"] = sub
    ws["A2"].font = SUB_FONT
    ws.row_dimensions[1].height = 30
    ws.sheet_view.showGridLines = False


def legend(ws, cell="A3"):
    ws[cell] = "Yellow cells = type here   ·   Grey cells = calculated (don't type)   ·   Slate headers = helper columns"
    ws[cell].font = Font(name=F, size=9, color=SLATE)


def dv_list(ws, formula, ref, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "Choose from the list", prompt, True
    dv.error, dv.errorTitle, dv.showErrorMessage = "Please pick a value from the list (edit lists on the Settings tab).", "Not in list", True
    ws.add_data_validation(dv)
    dv.add(ref)


def rng(col, rows):
    return f"${col}${rows[0]}:${col}${rows[1]}"


# ---------------------------------------------------------------- sheets

def build_settings(wb, today_override):
    ws = wb.create_sheet("Settings")
    title(ws, "Settings & Lists", "Edit your business details and dropdown lists here. Other tabs read from this sheet.")
    legend(ws)
    rows = [
        ("Your business name", "Sam Rivera Design (sample)", "in", None),
        ("Today (used by all formulas)", "=IF(C9<>\"\",C9,TODAY())", "calc", DATE_FMT),
        ("Default payment terms (days)", 14, "in", "0"),
        ("Currency code (for labels)", "USD", "in", None),
        ("Today override (leave blank)", today_override, "in", DATE_FMT),
        ("Weekly review day", "Friday", "in", None),
    ]
    ws["B4"], ws["C4"] = "Setting", "Value"
    for c in ("B4", "C4"):
        ws[c].font, ws[c].fill, ws[c].alignment, ws[c].border = H_FONT, HEAD_FILL, CENTER, BORDER
    for i, (label, val, kind, fmt) in enumerate(rows, start=5):
        ws.cell(row=i, column=2, value=label).font = BOLD
        c = ws.cell(row=i, column=3, value=val)
        c.font = BODY
        c.fill = INPUT_FILL if kind == "in" else CALC_FILL
        c.border = BORDER
        ws.cell(row=i, column=2).border = BORDER
        if fmt:
            c.number_format = fmt
    ws["C6"].comment = Comment("Uses TODAY() unless you type a date in 'Today override' (C9). The override is handy for testing or reviewing 'as of' a past date.", "Band of One")
    ws["C9"].comment = Comment("Leave blank for normal use.", "Band of One")
    ws["C10"].comment = Comment("Pick the day you'll do your weekly admin review.", "Band of One")
    dv_list(ws, '"' + ",".join(WEEKDAYS) + '"', "C10")
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 30

    # Lists block starting at row 13
    top = 13
    blocks = [
        ("E", "Client stages", [s[0] for s in STAGES]),
        ("F", "Win probability", [s[1] for s in STAGES]),
        ("G", "Stage group", [s[2] for s in STAGES]),
        ("I", "Admin decisions", DECISIONS),
        ("J", "Admin categories", CATEGORIES),
        ("K", "Project status", PROJ_STATUS),
        ("L", "Lead sources", SOURCES),
        ("M", "Services (edit me)", SERVICES + [""] * 8),
        ("O", "Reminder stage", [r[0] for r in REMINDERS]),
        ("P", "Days from due date", [r[1] for r in REMINDERS]),
        ("Q", "Meaning", [r[2] for r in REMINDERS]),
    ]
    ws[f"E{top-1}"] = "Dropdown lists (you can edit the text; keep the same number of stages)"
    ws[f"E{top-1}"].font = BOLD
    for col, head, values in blocks:
        h = ws[f"{col}{top}"]
        h.value, h.font, h.fill, h.alignment, h.border = head, H_FONT, HEAD_FILL, CENTER, BORDER
        ws.column_dimensions[col].width = 20 if col not in ("F", "P") else 14
        for i, v in enumerate(values, start=top + 1):
            c = ws[f"{col}{i}"]
            c.value = v
            c.font, c.fill, c.border = BODY, INPUT_FILL, BORDER
            if col == "F":
                c.number_format = "0%"
    ws.column_dimensions["Q"].width = 24
    ws.row_dimensions[top].height = 30
    ws[f"E{top+10}"] = "Win probability is an assumption used only to weight your pipeline on the Dashboard. Edit it to match your own experience."
    ws[f"E{top+10}"].font = Font(name=F, size=9, italic=True, color=SLATE)
    ws[f"O{top+7}"] = "Reminder timing follows the 5-stage cadence in The Solo Admin Playbook (Ch. 9). Change the days to suit your terms."
    ws[f"O{top+7}"].font = Font(name=F, size=9, italic=True, color=SLATE)
    return ws


def S(col, n):
    """Absolute range on Settings list block (rows 14..14+n-1)."""
    return f"Settings!${col}$14:${col}${13 + n}"


def build_clients(wb):
    ws = wb.create_sheet("Clients & Pipeline")
    title(ws, "Clients & Pipeline", "One row per client or lead. Update the Stage and Next action date and everything else calculates.")
    legend(ws)
    headers = ["Client ID", "Client / company", "Contact name", "Email", "Service", "Source", "Stage", "Project value",
               "Win %", "Weighted value", "Group", "First contact", "Proposal sent", "Next action", "Next action date",
               "Days until next action", "Action status", "Project end date", "90-day check-in", "Notes", "Active flag", "Action rank"]
    widths = [9, 24, 16, 24, 18, 14, 20, 13, 8, 13, 10, 12, 12, 30, 13, 11, 14, 12, 12, 30, 9, 9]
    kinds = ["calc", "in", "in", "in", "in", "in", "in", "in", "calc", "calc", "calc", "in", "in", "in", "in", "calc", "calc", "in", "calc", "in", "help", "help"]
    header(ws, 4, headers, widths, kinds)
    r0, r1 = CLIENT_ROWS
    n_st = len(STAGES)
    for r in range(r0, r1 + 1):
        ws[f"A{r}"] = f'=IF(B{r}="","","C"&TEXT(ROW()-4,"000"))'
        ws[f"I{r}"] = f'=IF(OR(B{r}="",G{r}=""),"",INDEX({S("F", n_st)},MATCH(G{r},{S("E", n_st)},0)))'
        ws[f"J{r}"] = f'=IF(OR(I{r}="",H{r}=""),"",H{r}*I{r})'
        ws[f"K{r}"] = f'=IF(OR(B{r}="",G{r}=""),"",INDEX({S("G", n_st)},MATCH(G{r},{S("E", n_st)},0)))'
        ws[f"P{r}"] = f'=IF(OR(B{r}="",O{r}="",G{r}="Lost"),"",O{r}-{TODAY})'
        ws[f"Q{r}"] = f'=IF(P{r}="","",IF(P{r}<0,"Overdue",IF(P{r}<=7,"Due this week","Upcoming")))'
        ws[f"S{r}"] = f'=IF(AND(G{r}="7 · Past client",R{r}<>""),R{r}+90,"")'
        ws[f"U{r}"] = f'=IF(P{r}="",0,1)'
        ws[f"V{r}"] = f'=IF(U{r}=0,"",COUNTIFS($U${r0}:$U${r1},1,$O${r0}:$O${r1},"<"&O{r})+COUNTIFS($U${r0}:$U{r},1,$O${r0}:$O{r},O{r}))'
    for col, kind in zip("ABCDEFGHIJKLMNOPQRSTUV", kinds):
        fill = INPUT_FILL if kind == "in" else (CALC_FILL if kind == "calc" else HELP_FILL)
        style_range(ws, f"{col}{r0}:{col}{r1}", font=BODY, fill=fill)
    for col in "LMORS":
        style_range(ws, f"{col}{r0}:{col}{r1}", fmt=DATE_FMT)
    for col in "HJ":
        style_range(ws, f"{col}{r0}:{col}{r1}", fmt=MONEY_FMT)
    style_range(ws, f"I{r0}:I{r1}", fmt="0%")
    style_range(ws, f"P{r0}:P{r1}", fmt="0")
    dv_list(ws, f"={S('E', n_st)}", f"G{r0}:G{r1}", "Pick the client's current stage.")
    dv_list(ws, f"={S('L', len(SOURCES))}", f"F{r0}:F{r1}")
    dv_list(ws, "=Settings!$M$14:$M$25", f"E{r0}:E{r1}")
    red = PatternFill("solid", fgColor="FBE3DA")
    yel = PatternFill("solid", fgColor="FDF1CC")
    ws.conditional_formatting.add(f"Q{r0}:Q{r1}", FormulaRule(formula=[f'Q{r0}="Overdue"'], fill=red, font=Font(name=F, bold=True, color=RUST)))
    ws.conditional_formatting.add(f"Q{r0}:Q{r1}", FormulaRule(formula=[f'Q{r0}="Due this week"'], fill=yel, font=Font(name=F, bold=True, color=INK)))
    ws.conditional_formatting.add(f"S{r0}:S{r1}", FormulaRule(formula=[f'AND(S{r0}<>"",S{r0}<=TODAY()+7)'], fill=yel))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:T{r1}"
    return ws


def build_projects(wb):
    ws = wb.create_sheet("Projects")
    title(ws, "Projects", "Track each project's milestones, deposit and weekly update day.")
    legend(ws)
    headers = ["Project", "Client", "Start date", "Due date", "Status", "Fee", "Deposit %", "Deposit amount", "Next milestone",
               "Milestone date", "Days to milestone", "Flag", "Weekly update day", "Notes"]
    widths = [26, 22, 12, 12, 16, 12, 10, 13, 28, 13, 11, 13, 13, 30]
    kinds = ["in", "in", "in", "in", "in", "in", "in", "calc", "in", "in", "calc", "calc", "in", "in"]
    header(ws, 4, headers, widths, kinds)
    r0, r1 = PROJECT_ROWS
    for r in range(r0, r1 + 1):
        ws[f"H{r}"] = f'=IF(OR(F{r}="",G{r}=""),"",F{r}*G{r})'
        ws[f"K{r}"] = f'=IF(OR(A{r}="",J{r}="",E{r}="Complete"),"",J{r}-{TODAY})'
        ws[f"L{r}"] = f'=IF(K{r}="","",IF(K{r}<0,"Overdue",IF(K{r}<=7,"This week","")))'
    for col, kind in zip("ABCDEFGHIJKLMN", kinds):
        style_range(ws, f"{col}{r0}:{col}{r1}", font=BODY, fill=INPUT_FILL if kind == "in" else CALC_FILL)
    for col in "CDJ":
        style_range(ws, f"{col}{r0}:{col}{r1}", fmt=DATE_FMT)
    for col in "FH":
        style_range(ws, f"{col}{r0}:{col}{r1}", fmt=MONEY_FMT)
    style_range(ws, f"G{r0}:G{r1}", fmt="0%")
    style_range(ws, f"K{r0}:K{r1}", fmt="0")
    dv_list(ws, f"={S('K', len(PROJ_STATUS))}", f"E{r0}:E{r1}")
    dv_list(ws, f"='Clients & Pipeline'!$B${CLIENT_ROWS[0]}:$B${CLIENT_ROWS[1]}", f"B{r0}:B{r1}", "Pick a client from the Clients & Pipeline tab.")
    dv_list(ws, '"' + ",".join(WEEKDAYS) + '"', f"M{r0}:M{r1}")
    ws.conditional_formatting.add(f"L{r0}:L{r1}", FormulaRule(formula=[f'L{r0}="Overdue"'], fill=PatternFill("solid", fgColor="FBE3DA"), font=Font(name=F, bold=True, color=RUST)))
    ws.conditional_formatting.add(f"L{r0}:L{r1}", FormulaRule(formula=[f'L{r0}="This week"'], fill=PatternFill("solid", fgColor="FDF1CC")))
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:N{r1}"
    return ws


def build_invoices(wb):
    ws = wb.create_sheet("Invoices & Follow-ups")
    title(ws, "Invoices & Follow-ups", "Log each invoice. Due dates, days overdue and the next reminder stage (A–E) calculate automatically.")
    legend(ws)
    headers = ["Invoice #", "Client", "Description", "Issue date", "Terms (days)", "Due date", "Amount", "Amount paid", "Paid date",
               "Balance", "Status", "Days overdue", "Reminder stage due", "Next reminder date", "Last reminder sent", "Last reminder date",
               "Action", "Notes", "Action flag", "Action rank"]
    widths = [11, 22, 26, 12, 9, 12, 12, 12, 12, 12, 12, 10, 11, 13, 11, 13, 16, 28, 9, 9]
    kinds = ["in", "in", "in", "in", "in", "calc", "in", "in", "in", "calc", "calc", "calc", "calc", "calc", "in", "in", "calc", "in", "help", "help"]
    header(ws, 4, headers, widths, kinds)
    r0, r1 = INVOICE_ROWS
    off = {r[0]: f"Settings!$P${14 + i}" for i, r in enumerate(REMINDERS)}
    for r in range(r0, r1 + 1):
        ws[f"F{r}"] = f'=IF(D{r}="","",D{r}+IF(E{r}="",{TERMS},E{r}))'
        ws[f"J{r}"] = f'=IF(G{r}="","",G{r}-N(H{r}))'
        ws[f"K{r}"] = (f'=IF(OR(G{r}="",F{r}=""),"",IF(J{r}<=0,"Paid",IF({TODAY}<F{r},IF(N(H{r})>0,"Part-paid","Not due"),'
                       f'IF({TODAY}=F{r},"Due today","Overdue"))))')
        ws[f"L{r}"] = f'=IF(OR(G{r}="",F{r}=""),"",IF(J{r}<=0,"",MAX(0,{TODAY}-F{r})))'
        ws[f"M{r}"] = (f'=IF(OR(G{r}="",F{r}=""),"",IF(J{r}<=0,"Paid",'
                       f'IF({TODAY}>=F{r}+{off["E"]},"E",IF({TODAY}>=F{r}+{off["D"]},"D",IF({TODAY}>=F{r}+{off["C"]},"C",'
                       f'IF({TODAY}>=F{r}+{off["B"]},"B",IF({TODAY}>=F{r}+{off["A"]},"A","—")))))))')
        ws[f"N{r}"] = (f'=IF(OR(G{r}="",F{r}="",J{r}<=0),"",IF({TODAY}<F{r}+{off["A"]},F{r}+{off["A"]},IF({TODAY}<F{r}+{off["B"]},F{r}+{off["B"]},'
                       f'IF({TODAY}<F{r}+{off["C"]},F{r}+{off["C"]},IF({TODAY}<F{r}+{off["D"]},F{r}+{off["D"]},IF({TODAY}<F{r}+{off["E"]},F{r}+{off["E"]},""))))))')
        ws[f"Q{r}"] = (f'=IF(OR(M{r}="",M{r}="Paid",M{r}="—"),IF(M{r}="Paid","Paid ✓",""),'
                       f'IF(O{r}="","Send "&M{r}&" now",IF(CODE(M{r})>CODE(O{r}),"Send "&M{r}&" now","Up to date")))')
        ws[f"S{r}"] = f'=IF(LEFT(Q{r},4)="Send",1,0)'
        ws[f"T{r}"] = f'=IF(S{r}=0,"",COUNTIFS($S${r0}:$S${r1},1,$F${r0}:$F${r1},"<"&F{r})+COUNTIFS($S${r0}:$S{r},1,$F${r0}:$F{r},F{r}))'
    for col, kind in zip("ABCDEFGHIJKLMNOPQRST", kinds):
        fill = INPUT_FILL if kind == "in" else (CALC_FILL if kind == "calc" else HELP_FILL)
        style_range(ws, f"{col}{r0}:{col}{r1}", font=BODY, fill=fill)
    for col in "DFINP":
        style_range(ws, f"{col}{r0}:{col}{r1}", fmt=DATE_FMT)
    for col in "GHJ":
        style_range(ws, f"{col}{r0}:{col}{r1}", fmt=MONEY_FMT)
    style_range(ws, f"L{r0}:L{r1}", fmt="0")
    style_range(ws, f"M{r0}:M{r1}", align=CENTER)
    dv_list(ws, f"='Clients & Pipeline'!$B${CLIENT_ROWS[0]}:$B${CLIENT_ROWS[1]}", f"B{r0}:B{r1}", "Pick a client from the Clients & Pipeline tab.")
    dv_list(ws, '"A,B,C,D,E"', f"O{r0}:O{r1}", "After you send a reminder, record its stage (A–E) and the date.")
    red = PatternFill("solid", fgColor="FBE3DA")
    ws.conditional_formatting.add(f"K{r0}:K{r1}", FormulaRule(formula=[f'K{r0}="Overdue"'], fill=red, font=Font(name=F, bold=True, color=RUST)))
    ws.conditional_formatting.add(f"K{r0}:K{r1}", FormulaRule(formula=[f'K{r0}="Due today"'], fill=PatternFill("solid", fgColor="FDF1CC"), font=Font(name=F, bold=True)))
    ws.conditional_formatting.add(f"K{r0}:K{r1}", FormulaRule(formula=[f'K{r0}="Paid"'], fill=PatternFill("solid", fgColor=MIST), font=Font(name=F, color=DEEP)))
    ws.conditional_formatting.add(f"Q{r0}:Q{r1}", FormulaRule(formula=[f'LEFT(Q{r0},4)="Send"'], fill=red, font=Font(name=F, bold=True, color=RUST)))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:R{r1}"
    return ws


def build_audit(wb):
    ws = wb.create_sheet("Admin Audit")
    title(ws, "Admin Audit", "List recurring admin tasks. Priority = (times/month × minutes) × repeatability ÷ risk. Fix the top-ranked tasks first.")
    legend(ws)
    headers = ["Task", "Category", "Times / month", "Minutes each", "Minutes / month", "Repeatability (1–3)", "Risk (1–3)",
               "Priority score", "Rank", "Decision", "Prompt / template ID", "Minutes / month after", "Minutes saved / month", "Notes"]
    widths = [30, 18, 10, 10, 11, 12, 10, 11, 7, 13, 14, 12, 12, 30]
    kinds = ["in", "in", "in", "in", "calc", "in", "in", "calc", "calc", "in", "in", "in", "calc", "in"]
    header(ws, 4, headers, widths, kinds)
    r0, r1 = AUDIT_ROWS
    for r in range(r0, r1 + 1):
        ws[f"E{r}"] = f'=IF(OR(C{r}="",D{r}=""),"",C{r}*D{r})'
        ws[f"H{r}"] = f'=IF(OR(E{r}="",F{r}="",G{r}=""),"",E{r}*F{r}/G{r})'
        ws[f"I{r}"] = f'=IF(H{r}="","",RANK(H{r},$H${r0}:$H${r1},0)+COUNTIF($H${r0}:H{r},H{r})-1)'
        ws[f"M{r}"] = f'=IF(OR(E{r}="",L{r}=""),"",E{r}-L{r})'
    for col, kind in zip("ABCDEFGHIJKLMN", kinds):
        style_range(ws, f"{col}{r0}:{col}{r1}", font=BODY, fill=INPUT_FILL if kind == "in" else CALC_FILL)
    style_range(ws, f"H{r0}:H{r1}", fmt="#,##0")
    style_range(ws, f"E{r0}:E{r1}", fmt="#,##0")
    style_range(ws, f"M{r0}:M{r1}", fmt="#,##0")
    dv_list(ws, f"={S('J', len(CATEGORIES))}", f"B{r0}:B{r1}")
    dv_list(ws, f"={S('I', len(DECISIONS))}", f"J{r0}:J{r1}")
    dv_list(ws, '"1,2,3"', f"F{r0}:F{r1}", "3 = nearly identical every time; 1 = different every time")
    dv_list(ws, '"1,2,3"', f"G{r0}:G{r1}", "1 = low stakes; 2 = money/client data/deadlines; 3 = legal, sensitive or relationship-critical")
    ws.conditional_formatting.add(f"A{r0}:I{r1}", FormulaRule(formula=[f'AND($I{r0}<>"",$I{r0}<=3)'], font=Font(name=F, bold=True, color=DEEP)))
    # summary block
    ws["P4"], ws["Q4"] = "Summary", "Value"
    for c in ("P4", "Q4"):
        ws[c].font, ws[c].fill, ws[c].alignment, ws[c].border = H_FONT, HEAD_FILL, CENTER, BORDER
    summary = [
        ("Admin minutes / month", f"=SUM(E{r0}:E{r1})", "#,##0"),
        ("Admin hours / month", "=Q5/60", "0.0"),
        ("Hours / month after changes (re-estimated tasks)", f'=IF(COUNT(M{r0}:M{r1})=0,"",(Q5-SUM(M{r0}:M{r1}))/60)', "0.0"),
        ("Hours saved / month (measured)", f'=IF(COUNT(M{r0}:M{r1})=0,"",SUM(M{r0}:M{r1})/60)', "0.0"),
        ("Tasks listed", f'=COUNTA(A{r0}:A{r1})', "0"),
    ]
    for i, (lab, f, fmt) in enumerate(summary, start=5):
        ws[f"P{i}"], ws[f"Q{i}"] = lab, f
        ws[f"P{i}"].font, ws[f"Q{i}"].font = BOLD, BODY
        ws[f"Q{i}"].fill, ws[f"Q{i}"].number_format = CALC_FILL, fmt
        ws[f"P{i}"].border = ws[f"Q{i}"].border = BORDER
        ws[f"P{i}"].alignment = WRAP
    ws.column_dimensions["O"].width = 3
    ws.column_dimensions["P"].width = 30
    ws.column_dimensions["Q"].width = 12
    ws["P11"] = "Savings are only shown once you re-estimate a task's minutes after changing how you do it. We never assume a saving."
    ws["P11"].font = Font(name=F, size=9, italic=True, color=SLATE)
    ws["P11"].alignment = WRAP
    ws.row_dimensions[11].height = 48
    ws.freeze_panes = "B5"
    return ws


def build_prompt_log(wb):
    ws = wb.create_sheet("Prompt Log")
    title(ws, "Prompt Log", "Keep track of the prompts you actually use. IDs match The Band of One Prompt Library (100 prompts).")
    legend(ws)
    headers = ["Prompt ID", "Prompt name", "Stage / category", "Where I saved it", "Times used", "Last used", "Rating (1–5)", "My tweaks / notes"]
    widths = [11, 38, 20, 24, 10, 12, 10, 40]
    kinds = ["in"] * 8
    header(ws, 4, headers, widths, kinds)
    r0, r1 = PROMPT_ROWS
    style_range(ws, f"A{r0}:H{r1}", font=BODY, fill=INPUT_FILL)
    style_range(ws, f"F{r0}:F{r1}", fmt=DATE_FMT)
    dv_list(ws, '"1,2,3,4,5"', f"G{r0}:G{r1}")
    ws.freeze_panes = "B5"
    return ws


def build_sops(wb):
    ws = wb.create_sheet("SOP Library")
    title(ws, "SOP Library", "Index of your standard operating procedures and when each is due for review.")
    legend(ws)
    headers = ["SOP name", "Trigger (what starts it)", "Where it lives (link)", "Last reviewed", "Review every (days)", "Next review", "Status", "Notes"]
    widths = [30, 30, 30, 13, 12, 13, 14, 30]
    kinds = ["in", "in", "in", "in", "in", "calc", "calc", "in"]
    header(ws, 4, headers, widths, kinds)
    r0, r1 = SOP_ROWS
    for r in range(r0, r1 + 1):
        ws[f"F{r}"] = f'=IF(OR(A{r}="",D{r}=""),"",D{r}+IF(E{r}="",90,E{r}))'
        ws[f"G{r}"] = f'=IF(A{r}="","",IF(D{r}="","Not written yet",IF(F{r}<={TODAY},"Review due","OK")))'
    for col, kind in zip("ABCDEFGH", kinds):
        style_range(ws, f"{col}{r0}:{col}{r1}", font=BODY, fill=INPUT_FILL if kind == "in" else CALC_FILL)
    style_range(ws, f"D{r0}:D{r1}", fmt=DATE_FMT)
    style_range(ws, f"F{r0}:F{r1}", fmt=DATE_FMT)
    ws.conditional_formatting.add(f"G{r0}:G{r1}", FormulaRule(formula=[f'G{r0}="Review due"'], fill=PatternFill("solid", fgColor="FDF1CC"), font=Font(name=F, bold=True)))
    ws.conditional_formatting.add(f"G{r0}:G{r1}", FormulaRule(formula=[f'G{r0}="Not written yet"'], font=Font(name=F, italic=True, color=SLATE)))
    ws.freeze_panes = "B5"
    return ws


def build_review(wb):
    ws = wb.create_sheet("Weekly Review")
    title(ws, "Weekly Review", "Your 45-minute weekly admin review: agenda, live numbers to copy, and a log to spot trends.")
    agenda = [
        ("1. Clear the decks (10 min)", "Process inbox & messages; capture tasks into your tracker."),
        ("2. Pipeline (10 min)", "Check every client's stage and next-action date; send due follow-ups."),
        ("3. Money (10 min)", "Invoices to send; reminders due (stages A–E); record payments."),
        ("4. Next week (10 min)", "Block time for client work and deadlines; prep update notes."),
        ("5. Improve one thing (5 min)", "Turn one annoying task into a template, SOP or prompt."),
    ]
    ws["A4"], ws["B4"] = "Agenda", "What to do"
    for c in ("A4", "B4"):
        ws[c].font, ws[c].fill, ws[c].border = H_FONT, HEAD_FILL, BORDER
    for i, (a, b) in enumerate(agenda, start=5):
        ws[f"A{i}"], ws[f"B{i}"] = a, b
        ws[f"A{i}"].font, ws[f"B{i}"].font = BOLD, BODY
        ws[f"A{i}"].border = ws[f"B{i}"].border = BORDER
    # live numbers
    ws["E4"], ws["F4"] = "Live numbers (copy into the log)", "Now"
    for c in ("E4", "F4"):
        ws[c].font, ws[c].fill, ws[c].border = H_FONT, HEAD_FILL, BORDER
    i0, i1 = INVOICE_ROWS
    c0, c1 = CLIENT_ROWS
    live = [
        ("Overdue invoices (count)", f"=COUNTIF('Invoices & Follow-ups'!$K${i0}:$K${i1},\"Overdue\")", "0"),
        ("Overdue invoices (total)", f"=SUMIFS('Invoices & Follow-ups'!$J${i0}:$J${i1},'Invoices & Follow-ups'!$K${i0}:$K${i1},\"Overdue\")", MONEY_FMT),
        ("Reminders to send now", f"=SUM('Invoices & Follow-ups'!$S${i0}:$S${i1})", "0"),
        ("Client follow-ups overdue", f"=COUNTIF('Clients & Pipeline'!$Q${c0}:$Q${c1},\"Overdue\")", "0"),
        ("Client follow-ups due this week", f"=COUNTIF('Clients & Pipeline'!$Q${c0}:$Q${c1},\"Due this week\")", "0"),
    ]
    for i, (lab, f, fmt) in enumerate(live, start=5):
        ws[f"E{i}"], ws[f"F{i}"] = lab, f
        ws[f"E{i}"].font, ws[f"F{i}"].font = BOLD, BODY
        ws[f"F{i}"].fill, ws[f"F{i}"].number_format = CALC_FILL, fmt
        ws[f"E{i}"].border = ws[f"F{i}"].border = BORDER
    headers = ["Week starting", "Admin hours (est.)", "Overdue invoices (#)", "Overdue total", "Avg first reply (hours)",
               "Proposals sent", "Proposals won", "Win rate", "Improvement I made", "Notes"]
    widths = [30, 60, 12, 14, 34, 12, 12, 10, 34, 30]
    kinds = ["in", "in", "in", "in", "in", "in", "in", "calc", "in", "in"]
    hr = REVIEW_ROWS[0] - 1
    ws[f"A{hr-1}"] = "Weekly log"
    ws[f"A{hr-1}"].font = Font(name=F, bold=True, size=12, color=DEEP)
    header(ws, hr, headers, widths, kinds)
    r0, r1 = REVIEW_ROWS
    for r in range(r0, r1 + 1):
        ws[f"H{r}"] = f'=IF(OR(F{r}="",F{r}=0),"",G{r}/F{r})'
    for col, kind in zip("ABCDEFGHIJ", kinds):
        style_range(ws, f"{col}{r0}:{col}{r1}", font=BODY, fill=INPUT_FILL if kind == "in" else CALC_FILL)
    style_range(ws, f"A{r0}:A{r1}", fmt=DATE_FMT)
    style_range(ws, f"D{r0}:D{r1}", fmt=MONEY_FMT)
    style_range(ws, f"H{r0}:H{r1}", fmt="0%")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["E"].width = 34
    ws.column_dimensions["F"].width = 12
    ws.freeze_panes = f"B{r0}"
    return ws


def build_dashboard(wb):
    ws = wb.create_sheet("Dashboard", 1)
    ws.sheet_view.showGridLines = False
    ws["B1"] = "The Client Desk"
    ws["B1"].font = TITLE_FONT
    ws["B2"] = '="Dashboard for "&Settings!$C$5&" · as of "&TEXT(Settings!$C$6,"d mmm yyyy")'
    ws["B2"].font = SUB_FONT
    for col, w in zip("ABCDEFGHIJKLM", [2, 24, 16, 14, 13, 3, 24, 16, 14, 13, 3, 20, 14]):
        ws.column_dimensions[col].width = w
    c0, c1 = CLIENT_ROWS
    i0, i1 = INVOICE_ROWS
    a0, a1 = AUDIT_ROWS
    CL, INV, AUD = "'Clients & Pipeline'", "'Invoices & Follow-ups'", "'Admin Audit'"
    cur = "Settings!$C$8"
    tiles = [
        ("B4", "Active clients", f'=COUNTIF({CL}!$K${c0}:$K${c1},"Active")', "0"),
        ("D4", '="Weighted pipeline ("&' + cur + '&")"', f'=SUMIFS({CL}!$J${c0}:$J${c1},{CL}!$K${c0}:$K${c1},"Pipeline")', "#,##0"),
        ("G4", '="Outstanding ("&' + cur + '&")"', f'=SUMIFS({INV}!$J${i0}:$J${i1},{INV}!$K${i0}:$K${i1},"<>Paid")', "#,##0"),
        ("I4", '="Overdue ("&' + cur + '&")"', f'=SUMIFS({INV}!$J${i0}:$J${i1},{INV}!$K${i0}:$K${i1},"Overdue")', "#,##0"),
        ("L4", "Reminders to send", f"=SUM({INV}!$S${i0}:$S${i1})", "0"),
        ("B8", "Follow-ups overdue", f'=COUNTIF({CL}!$Q${c0}:$Q${c1},"Overdue")', "0"),
        ("D8", "Follow-ups this week", f'=COUNTIF({CL}!$Q${c0}:$Q${c1},"Due this week")', "0"),
        ("G8", '="Invoiced this month ("&' + cur + '&")"', f'=SUMIFS({INV}!$G${i0}:$G${i1},{INV}!$D${i0}:$D${i1},">="&DATE(YEAR(Settings!$C$6),MONTH(Settings!$C$6),1),{INV}!$D${i0}:$D${i1},"<="&Settings!$C$6)', "#,##0"),
        ("I8", '="Collected this month ("&' + cur + '&")"', f'=SUMIFS({INV}!$H${i0}:$H${i1},{INV}!$I${i0}:$I${i1},">="&DATE(YEAR(Settings!$C$6),MONTH(Settings!$C$6),1),{INV}!$I${i0}:$I${i1},"<="&Settings!$C$6)', "#,##0"),
        ("L8", "Admin hours / month", f"={AUD}!$Q$6", "0.0"),
    ]
    for anchor, label, formula, fmt in tiles:
        col = anchor[0]
        row = int(anchor[1:])
        nxt = get_column_letter(ws[anchor].column + 1)
        ws.merge_cells(f"{col}{row}:{nxt}{row}")
        ws.merge_cells(f"{col}{row+1}:{nxt}{row+2}")
        ws[f"{col}{row}"] = label
        ws[f"{col}{row}"].font = Font(name=F, size=9, bold=True, color=SLATE)
        ws[f"{col}{row+1}"] = formula
        ws[f"{col}{row+1}"].font = Font(name=F, size=20, bold=True, color=DEEP)
        ws[f"{col}{row+1}"].number_format = fmt
        ws[f"{col}{row+1}"].alignment = Alignment(horizontal="left", vertical="center")
        for r in range(row, row + 3):
            for cc in (col, nxt):
                ws[f"{cc}{r}"].fill = TILE_FILL
    # Next client actions
    def section(cell, text):
        ws[cell] = text
        ws[cell].font = Font(name=F, size=12, bold=True, color=DEEP)

    section("B12", "Next client actions")
    heads = ["Client", "Next action", "Date", "Status"]
    for j, h in enumerate(heads):
        c = ws.cell(row=13, column=2 + j, value=h)
        c.font, c.fill, c.border, c.alignment = H_FONT, HEAD_FILL, BORDER, CENTER
    for k in range(1, 9):
        r = 13 + k
        m = f"MATCH({k},{CL}!$V${c0}:$V${c1},0)"
        ws[f"B{r}"] = f'=IFERROR(INDEX({CL}!$B${c0}:$B${c1},{m}),"")'
        ws[f"C{r}"] = f'=IFERROR(INDEX({CL}!$N${c0}:$N${c1},{m}),"")'
        ws[f"D{r}"] = f'=IFERROR(INDEX({CL}!$O${c0}:$O${c1},{m}),"")'
        ws[f"E{r}"] = f'=IFERROR(INDEX({CL}!$Q${c0}:$Q${c1},{m}),"")'
        style_range(ws, f"B{r}:E{r}", font=BODY, fill=WHITE)
        ws[f"D{r}"].number_format = DATE_FMT
    ws.conditional_formatting.add("E14:E21", FormulaRule(formula=['E14="Overdue"'], font=Font(name=F, bold=True, color=RUST)))
    # Invoice actions
    section("G12", "Invoices needing a reminder")
    heads = ["Client", "Invoice #", "Balance", "Action"]
    for j, h in enumerate(heads):
        c = ws.cell(row=13, column=7 + j, value=h)
        c.font, c.fill, c.border, c.alignment = H_FONT, HEAD_FILL, BORDER, CENTER
    for k in range(1, 9):
        r = 13 + k
        m = f"MATCH({k},{INV}!$T${i0}:$T${i1},0)"
        ws[f"G{r}"] = f'=IFERROR(INDEX({INV}!$B${i0}:$B${i1},{m}),"")'
        ws[f"H{r}"] = f'=IFERROR(INDEX({INV}!$A${i0}:$A${i1},{m}),"")'
        ws[f"I{r}"] = f'=IFERROR(INDEX({INV}!$J${i0}:$J${i1},{m}),"")'
        ws[f"J{r}"] = f'=IFERROR(INDEX({INV}!$Q${i0}:$Q${i1},{m}),"")'
        style_range(ws, f"G{r}:J{r}", font=BODY, fill=WHITE)
        ws[f"I{r}"].number_format = MONEY_FMT
    ws.conditional_formatting.add("J14:J21", FormulaRule(formula=['LEFT(J14,4)="Send"'], font=Font(name=F, bold=True, color=RUST)))
    # Top admin tasks
    section("B24", "Top admin tasks to fix first")
    heads = ["Task", "Minutes / month", "Priority", "Decision"]
    for j, h in enumerate(heads):
        c = ws.cell(row=25, column=2 + j, value=h)
        c.font, c.fill, c.border, c.alignment = H_FONT, HEAD_FILL, BORDER, CENTER
    for k in range(1, 4):
        r = 25 + k
        m = f"MATCH({k},{AUD}!$I${a0}:$I${a1},0)"
        ws[f"B{r}"] = f'=IFERROR(INDEX({AUD}!$A${a0}:$A${a1},{m}),"")'
        ws[f"C{r}"] = f'=IFERROR(INDEX({AUD}!$E${a0}:$E${a1},{m}),"")'
        ws[f"D{r}"] = f'=IFERROR(INDEX({AUD}!$H${a0}:$H${a1},{m}),"")'
        ws[f"E{r}"] = f'=IFERROR(INDEX({AUD}!$J${a0}:$J${a1},{m}),"")'
        style_range(ws, f"B{r}:E{r}", font=BODY, fill=WHITE)
        ws[f"D{r}"].number_format = "#,##0"
    # Pipeline by stage
    section("G24", "Clients by stage")
    heads = ["Stage", "Clients", "Value"]
    for j, h in enumerate(heads):
        c = ws.cell(row=25, column=7 + j, value=h)
        c.font, c.fill, c.border, c.alignment = H_FONT, HEAD_FILL, BORDER, CENTER
    for k, (st, _, _) in enumerate(STAGES):
        r = 26 + k
        ws[f"G{r}"] = f"=Settings!$E${14 + k}"
        ws[f"H{r}"] = f'=COUNTIF({CL}!$G${c0}:$G${c1},G{r})'
        ws[f"I{r}"] = f'=SUMIFS({CL}!$H${c0}:$H${c1},{CL}!$G${c0}:$G${c1},G{r})'
        style_range(ws, f"G{r}:I{r}", font=BODY, fill=WHITE)
        ws[f"I{r}"].number_format = "#,##0"
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Clients by stage"
    chart.y_axis.title = None
    chart.x_axis.title = None
    data = Reference(ws, min_col=8, min_row=25, max_row=25 + len(STAGES))
    cats = Reference(ws, min_col=7, min_row=26, max_row=25 + len(STAGES))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.height, chart.width = 7.5, 14
    chart.series[0].graphicalProperties.solidFill = TEAL
    chart.x_axis.scaling.orientation = "maxMin"
    chart.y_axis.majorUnit = 1
    chart.y_axis.number_format = "0"
    chart.y_axis.delete = False
    chart.x_axis.delete = False
    ws.add_chart(chart, "B31")
    ws["G35"] = "Tip: sort or filter the other tabs any time. The dashboard reads the whole range."
    ws["G35"].font = Font(name=F, size=9, italic=True, color=SLATE)
    return ws


def build_start(wb):
    ws = wb.active
    ws.title = "Start Here"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 90
    ws["B1"] = "The Client Desk"
    ws["B1"].font = Font(name=F, bold=True, size=22, color=DEEP)
    ws["B2"] = "A simple command center for a business of one · by Band of One · v1.0"
    ws["B2"].font = SUB_FONT
    rows = [
        ("What this is", "One workbook for clients, pipeline, projects, invoices and payment reminders, plus the admin audit, prompt log, SOP index and weekly review from The Solo Admin Playbook."),
        ("Works in", "Microsoft Excel 2010 or later · Google Sheets (upload, then File → Save as Google Sheets) · LibreOffice Calc. No macros, no add-ins."),
        ("Colour key", "YELLOW cells = type here. GREY cells = calculated, don't type over them. SLATE headers = helper columns the dashboard needs (leave them)."),
        ("Step 1 · Settings", "Open the Settings tab. Enter your business name, default payment terms and currency code. Edit the Services list."),
        ("Step 2 · Clear the sample", "Rows marked SAMPLE in the Notes column are a fictional example business. Select their YELLOW cells and press Delete. Don't delete whole rows, so the formulas stay in place."),
        ("Step 3 · Add clients", "Clients & Pipeline: one row per client or lead. Choose a Stage from the dropdown and always fill in 'Next action' and 'Next action date'."),
        ("Step 4 · Log invoices", "Invoices & Follow-ups: enter issue date, amount and (optionally) terms. Due date, status, days overdue and the next reminder stage (A–E) calculate automatically. When you send a reminder, record its stage and date."),
        ("Step 5 · Run your audit", "Admin Audit: list recurring tasks, then score repeatability and risk. The highest-ranked tasks appear on the Dashboard. Re-estimate minutes after you change a task to measure real savings."),
        ("Weekly habit", "Every week (Settings → review day) open the Dashboard, work through 'Next client actions' and 'Invoices needing a reminder', then log your numbers on the Weekly Review tab."),
        ("Reminder stages", "A = heads-up 3 days before due · B = due date · C = 7 days overdue · D = 14 days overdue · E = 30 days overdue (final notice). Change the timing on Settings. Templates are in The Solo Admin Playbook, Chapter 9."),
        ("Dates", "Type dates as YYYY-MM-DD (e.g., 2026-10-15) or in your system's date format. 'Today' comes from your computer clock unless you set an override on Settings."),
        ("Capacity", "Clients: 150 rows · Invoices: 200 rows · Projects: 100 rows · Audit: 40 rows. To add more, copy the last row down (formulas copy with it)."),
        ("Privacy", "This file lives on your computer or your own cloud drive. Nothing is sent anywhere. Keep it private: it contains client details."),
        ("Not advice", "The workbook helps you organize information. It isn't accounting, tax or legal software. Keep your official records in the tools your accountant recommends."),
        ("License", "Personal business use by the purchaser. Don't resell or redistribute the file. See LICENSE.txt in your download."),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws[f"B{i}"], ws[f"C{i}"] = k, v
        ws[f"B{i}"].font = Font(name=F, bold=True, size=10, color=DEEP)
        ws[f"C{i}"].font = BODY
        ws[f"C{i}"].alignment = WRAP
        ws[f"B{i}"].alignment = WRAP
        ws.row_dimensions[i].height = 44 if len(v) > 150 else 30
        for c in (f"B{i}", f"C{i}"):
            ws[c].border = BORDER
            ws[c].fill = INPUT_FILL if k == "Colour key" else WHITE
    return ws


# ---------------------------------------------------------------- sample data

def sample(wb):
    T = "Settings!$C$6"
    cl = wb["Clients & Pipeline"]
    clients = [
        ("Harbor Dental Studio", "Dr. Lee (fictional)", "office@example.com", "Website design", "Referral", "5 · Delivery", 4800, f"={T}-60", f"={T}-45", "Send weekly update", f"={T}+2", "", "SAMPLE"),
        ("Oak & Pine Law", "J. Patel (fictional)", "jp@example.com", "Website design", "LinkedIn", "3 · Proposal sent", 6500, f"={T}-12", f"={T}-5", "Proposal follow-up (day 7)", f"={T}+2", "", "SAMPLE"),
        ("Brightside Physio", "M. Chen (fictional)", "hello@example.com", "Website care plan", "Website", "2 · Discovery", 1200, f"={T}-4", "", "Send call summary", f"={T}-1", "", "SAMPLE"),
        ("Northfield Accounting", "R. Gomez (fictional)", "rg@example.com", "Website design", "Past client", "4 · Onboarding", 5200, f"={T}-20", f"={T}-14", "Kickoff call prep", f"={T}+5", "", "SAMPLE"),
        ("Cedar Yoga", "A. Novak (fictional)", "studio@example.com", "Strategy session", "Instagram", "1 · Inquiry", 400, f"={T}-1", "", "Reply to inquiry", f"={T}", "", "SAMPLE"),
        ("Maple Street Bakery", "T. Brown (fictional)", "tb@example.com", "Website design", "Referral", "7 · Past client", 3900, f"={T}-200", f"={T}-190", "90-day check-in email", f"={T}+10", f"={T}-80", "SAMPLE"),
        ("Summit Coaching", "K. Ito (fictional)", "ki@example.com", "Website design", "Website", "Lost", 3000, f"={T}-90", f"={T}-80", "", "", "", "SAMPLE: chose a cheaper option"),
    ]
    cols = "BCDEFGHLMNORT"
    for i, row in enumerate(clients, start=CLIENT_ROWS[0]):
        for col, val in zip(cols, row):
            if val != "":
                cl[f"{col}{i}"] = val
    pj = wb["Projects"]
    projects = [
        ("Harbor Dental website", "Harbor Dental Studio", f"={T}-40", f"={T}+20", "In progress", 4800, 0.5, "Homepage design review", f"={T}+3", "Friday", "SAMPLE"),
        ("Northfield site rebuild", "Northfield Accounting", f"={T}+2", f"={T}+50", "Not started", 5200, 0.5, "Kickoff call", f"={T}+5", "Thursday", "SAMPLE"),
        ("Maple Street care plan", "Maple Street Bakery", f"={T}-80", f"={T}-10", "Complete", 3900, 0.5, "", "", "", "SAMPLE"),
    ]
    for i, row in enumerate(projects, start=PROJECT_ROWS[0]):
        for col, val in zip("ABCDEFGIJMN", row):
            if val != "":
                pj[f"{col}{i}"] = val
    inv = wb["Invoices & Follow-ups"]
    invoices = [
        ("2026-014", "Harbor Dental Studio", "Deposit 50%", f"={T}-45", 14, 2400, 2400, f"={T}-33", "", "", "SAMPLE: paid"),
        ("2026-018", "Harbor Dental Studio", "Milestone 2: design approved", f"={T}-22", 14, 1200, "", "", "C", f"={T}-1", "SAMPLE: 8 days overdue, stage C sent"),
        ("2026-019", "Maple Street Bakery", "Care plan: quarter", f"={T}-44", 14, 450, "", "", "C", f"={T}-23", "SAMPLE: 30 days overdue → final notice due"),
        ("2026-020", "Northfield Accounting", "Deposit 50%", f"={T}-12", 14, 2600, "", "", "", "", "SAMPLE: due in 2 days → heads-up"),
        ("2026-021", "Oak & Pine Law", "Strategy workshop", f"={T}-14", 14, 600, "", "", "A", f"={T}-3", "SAMPLE: due today"),
        ("2026-022", "Brightside Physio", "Audit call", f"={T}-2", 14, 250, 100, f"={T}-1", "", "", "SAMPLE: part-paid, not due"),
        ("2026-023", "Harbor Dental Studio", "Hosting setup", f"={T}-29", 14, 180, "", "", "", "", "SAMPLE: 15 days overdue, no reminder logged"),
    ]
    for i, row in enumerate(invoices, start=INVOICE_ROWS[0]):
        for col, val in zip("ABCDEGHIOPR", row):
            if val != "":
                inv[f"{col}{i}"] = val
    au = wb["Admin Audit"]
    audit = [
        ("Weekly client updates", "During projects", 16, 20, 3, 1, "Template", "DEL-01", "", "SAMPLE (Sam's audit)"),
        ("Social media posts", "Running the business", 8, 30, 2, 1, "AI draft", "MKT-01", "", "SAMPLE"),
        ("Replying to inquiries", "Getting clients", 8, 15, 3, 1, "Template", "LEAD-01", "", "SAMPLE"),
        ("Onboarding emails & setup", "Starting projects", 2, 60, 3, 1, "Template", "ONB-01", "", "SAMPLE"),
        ("Writing proposals", "Getting clients", 3, 90, 2, 2, "AI draft", "PROP-01", "", "SAMPLE"),
        ("Meeting notes & follow-ups", "During projects", 12, 15, 3, 2, "AI draft", "DEL-04", "", "SAMPLE"),
        ("Sending invoices", "Money", 6, 15, 3, 2, "Template", "", "", "SAMPLE"),
        ("Chasing late payments", "Money", 3, 20, 3, 2, "Template", "PAY-02", "", "SAMPLE"),
    ]
    for i, row in enumerate(audit, start=AUDIT_ROWS[0]):
        for col, val in zip("ABCDFGJKLN", row):
            if val != "":
                au[f"{col}{i}"] = val
    pl = wb["Prompt Log"]
    prompts = [
        ("SET-01", "Business brief builder", "Setup"), ("SET-02", "Voice profile from writing samples", "Setup"),
        ("SET-03", "Redaction check before pasting", "Setup"), ("LEAD-01", "Reply to a new inquiry", "Leads & inquiries"),
        ("DISC-02", "Discovery call prep sheet", "Discovery"), ("DISC-05", "Call notes → summary + next steps email", "Discovery"),
        ("PROP-01", "Proposal draft from discovery notes", "Proposals"), ("PROP-06", "Scope exclusions & assumptions list", "Proposals"),
        ("ONB-01", "Welcome email", "Onboarding"), ("ONB-03", "Intake questionnaire builder", "Onboarding"),
        ("DEL-01", "Weekly client update", "Delivery"), ("DEL-04", "Meeting notes → summary & action items", "Delivery"),
        ("DEL-07", "Scope change request email", "Delivery"), ("PAY-02", "Payment reminder sequence personalizer", "Getting paid"),
        ("RET-02", "Testimonial request with specific questions", "Offboarding"), ("OPS-01", "SOP from rough notes", "Systems"),
        ("OPS-05", "Weekly admin review assistant", "Systems"),
    ]
    for i, (pid, name, cat) in enumerate(prompts, start=PROMPT_ROWS[0]):
        pl[f"A{i}"], pl[f"B{i}"], pl[f"C{i}"] = pid, name, cat
    sp = wb["SOP Library"]
    sops = [
        ("Onboard a new client", "Client accepts proposal", "", f"={T}-30", 90), ("Send the weekly update", "Every Friday 3 pm", "", f"={T}-95", 90),
        ("Invoice and follow up", "Milestone approved / invoice due", "", "", 90), ("Offboard a client", "Final deliverable approved", "", "", 180),
        ("Weekly admin review", "Weekly review day", "", f"={T}-7", 90),
    ]
    for i, row in enumerate(sops, start=SOP_ROWS[0]):
        for col, val in zip("ABCDE", row):
            if val != "":
                sp[f"{col}{i}"] = val
        sp[f"H{i}"] = "SAMPLE"
    wr = wb["Weekly Review"]
    for i, row in enumerate([(f"={T}-14", 6.5, 2, 1650, 20, 2, 1, "Saved inquiry reply template", "SAMPLE"),
                             (f"={T}-7", 5.5, 1, 1200, 8, 1, 1, "Wrote onboarding SOP", "SAMPLE")], start=REVIEW_ROWS[0]):
        for col, val in zip("ABCDEFGIJ", row):
            wr[f"{col}{i}"] = val


def add_names(wb):
    names = {"TodayDate": "Settings!$C$6", "DefaultTerms": "Settings!$C$7", "CurrencyCode": "Settings!$C$8"}
    for n, ref in names.items():
        dn = DefinedName(n, attr_text=ref)
        wb.defined_names[n] = dn


def build(today_override=None, out=None):
    wb = Workbook()
    build_start(wb)
    build_settings(wb, today_override)
    build_clients(wb)
    build_projects(wb)
    build_invoices(wb)
    build_audit(wb)
    build_prompt_log(wb)
    build_sops(wb)
    build_review(wb)
    build_dashboard(wb)
    sample(wb)
    add_names(wb)
    order = ["Start Here", "Dashboard", "Clients & Pipeline", "Projects", "Invoices & Follow-ups", "Admin Audit", "Prompt Log", "SOP Library", "Weekly Review", "Settings"]
    wb._sheets = [wb[n] for n in order]
    colors = {"Start Here": BUTTER, "Dashboard": TEAL, "Settings": SLATE}
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = colors.get(ws.title, DEEP)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.sheet_view.zoomScale = 100
    wb.active = 0
    wb.properties.title = "The Client Desk"
    wb.properties.creator = "Band of One"
    out = Path(out) if out else OUT_DIR / FILE_NAME
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--today", help="YYYY-MM-DD override written into Settings!C9 (testing)")
    ap.add_argument("--out")
    a = ap.parse_args()
    t = dt.date.fromisoformat(a.today) if a.today else None
    print(build(t, a.out))
