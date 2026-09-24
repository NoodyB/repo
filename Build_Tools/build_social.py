"""Build social media deliverables for both businesses from the YAML sources.

Sources:
  Content_Automation_Business/05_Social_Media_Content/source/{videos,posts}.yaml        (Month 1: content)
  Digital_Products_Business/08_Marketing_Campaign/source/{launch_videos,launch_posts}.yaml (Month 2: launch)
Outputs: Markdown docs, scheduler CSVs, a 60-day calendar workbook (formula dates), graphics
(carousel slides 1080x1350, promo cards 1080x1350, video covers 1080x1920).

Run: python3 Build_Tools/build_social.py [--no-graphics]
"""
import csv
import html
import shutil
import sys

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from brand import COLORS as C, ROOT, font_face_css
from render import run, write_html

CB = ROOT / "Content_Automation_Business" / "05_Social_Media_Content"
PB = ROOT / "Digital_Products_Business" / "08_Marketing_Campaign"
LOGO_DARK = (ROOT / "Digital_Products_Business/02_Branding/logo/band-of-one-horizontal-on-dark.svg").as_uri()
LOGO = (ROOT / "Digital_Products_Business/02_Branding/logo/band-of-one-horizontal.svg").as_uri()
COVERS = {k: (ROOT / "Content_Automation_Business/03_Website/src/assets/img" / f"cover-{v}.webp").as_uri()
          for k, v in {"playbook": "playbook", "desk": "desk", "prompts": "prompts", "starter": "starter", "bundle": "prompts"}.items()}
PRODUCT_TAGS = {"prompts": "ai", "desk": "paid", "playbook": "systems", "bundle": "brand", "starter": "brand"}
RING = {"workflows": C["orange"], "paid": C["butter"], "ai": C["mist"], "systems": "#8FC7BE", "tools": C["rust"], "brand": C["orange"]}


def load(p):
    return yaml.safe_load(open(p, encoding="utf-8"))


# ------------------------------------------------------------------ documents

def video_doc(videos, title, intro):
    out = [f"# {title}", "", intro, ""]
    out += ["| # | Day | Concept | Hook | Link |", "|---|---|---|---|---|"]
    out += [f"| {v['id']} | {v['day']} | {v['concept']} | {v['hook']} | `{v['link']}` |" for v in videos]
    out.append("")
    for v in videos:
        out += [f"## {v['id']} · Day {v['day']} · {v['concept']}", "",
                f"**Hook (first 2 seconds):** {v['hook']}  ",
                f"**On-screen text:** {v['on_screen']}  ",
                f"**CTA:** {v['cta']} → `{v['link']}`", "",
                "**Script** (about 30–60 seconds):", "", "```text", v["script"].strip(), "```", "",
                "**Caption:**", "", f"> {v['caption']}", "",
                f"**Hashtags:** {' '.join(TAGS[v['tags']])}", ""]
    return "\n".join(out)


def posts_doc(posts, title, intro):
    out = [f"# {title}", "", intro, "", "## Educational posts", ""]
    for p in posts["educational"]:
        out += [f"### {p['id']} · Day {p['day']} · {p['platform']} ({p['format']}) · {p['title']}", ""]
        if p.get("slides"):
            out += ["**Slides:**", ""] + [f"{i}. {s}" for i, s in enumerate(p["slides"], 1)] + ["", "**Caption:**", ""]
        out += [p["text"].strip(), "", f"**Hashtags:** {' '.join(TAGS[post_tag(p)])}", ""]
    out += ["## Promotional posts", ""]
    for p in posts["promotional"]:
        out += [f"### {p['id']} · Day {p['day']} · {p['platform']} · {p['title']}", "", p["text"].strip(), "",
                f"**Hashtags:** {' '.join(TAGS[post_tag(p)])}", ""]
    return "\n".join(out)


def post_tag(p):
    return p.get("tags") or PRODUCT_TAGS.get(p.get("product"), "brand")


def calendar_rows(month, offset, videos, posts):
    rows = []
    for v in videos:
        rows.append(dict(day=v["day"] + offset, month=month, channel="TikTok / Reels / Shorts (+ LinkedIn native)",
                         type="Short video", id=v["id"], title=v["concept"], hook=v["hook"], link=v["link"],
                         cta=v["cta"], text=v["caption"], tags=" ".join(TAGS[v["tags"]]), asset=f"graphics/video_covers/{v['id']}.png"))
    for kind, items in (("Educational post", posts["educational"]), ("Promotional post", posts["promotional"])):
        for p in items:
            asset = f"graphics/carousels/{p['id']}_01.png" if p.get("slides") else (f"graphics/promo_cards/{p['id']}.png" if kind.startswith("Promo") else "")
            rows.append(dict(day=p["day"] + offset, month=month, channel=p["platform"], type=kind + (" (carousel)" if p.get("slides") else ""),
                             id=p["id"], title=p["title"], hook="", link=p.get("link", ""), cta="",
                             text=p["text"].strip(), tags=" ".join(TAGS[post_tag(p)]), asset=asset))
    return sorted(rows, key=lambda r: (r["day"], r["type"], r["id"]))


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Day", "Campaign month", "Channel", "Type", "ID", "Title / concept", "Hook", "Link", "CTA", "Caption / post text", "Hashtags", "Graphic", "Status"])
        for r in rows:
            w.writerow([r["day"], r["month"], r["channel"], r["type"], r["id"], r["title"], r["hook"], r["link"], r["cta"], r["text"], r["tags"], r["asset"], "Draft"])


def calendar_xlsx(rows, path):
    F = "Arial"
    thin = Side(style="thin", color="D9D4CA")
    B = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = "Calendar"
    ws["A1"] = "Band of One: 60-day social calendar (Month 1 content → Month 2 product launch)"
    ws["A1"].font = Font(name=F, bold=True, size=14, color=C["deep_teal"][1:])
    ws["A2"], ws["C2"] = "Start date (Day 1):", None
    ws["A2"].font = Font(name=F, bold=True)
    ws["C2"].fill = PatternFill("solid", fgColor="FFF8E1")
    ws["C2"].number_format = "yyyy-mm-dd"
    ws["E2"] = "Type your Day 1 date in the yellow cell; every date below updates. Launch day = Day 38 (Month 2, day 8). Nothing is scheduled or posted automatically."
    ws["E2"].font = Font(name=F, italic=True, size=9, color=C["slate"][1:])
    heads = ["Day", "Date", "Weekday", "Campaign month", "Channel", "Type", "ID", "Title / concept", "Hook", "Link", "Caption / post text", "Hashtags", "Graphic", "Status", "Posted URL", "Notes"]
    widths = [6, 12, 10, 16, 22, 20, 8, 30, 34, 30, 60, 34, 30, 12, 26, 20]
    for j, (h, w) in enumerate(zip(heads, widths)):
        c = ws.cell(row=4, column=1 + j, value=h)
        c.font, c.fill, c.border = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor=C["deep_teal"][1:]), B
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    for k, r in enumerate(rows):
        i = 5 + k
        vals = [r["day"], f'=IF($C$2="","",$C$2+A{i}-1)', f'=IF(B{i}="","",TEXT(B{i},"ddd"))', r["month"], r["channel"], r["type"], r["id"],
                r["title"], r["hook"], r["link"], r["text"], r["tags"], r["asset"], "Draft", "", ""]
        for j, v in enumerate(vals):
            c = ws.cell(row=i, column=1 + j, value=v)
            c.font = Font(name=F, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = B
        ws.cell(row=i, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=14).fill = PatternFill("solid", fgColor="FFF8E1")
        ws.cell(row=i, column=15).fill = PatternFill("solid", fgColor="FFF8E1")
    last = 4 + len(rows)
    dv = DataValidation(type="list", formula1='"Draft,Approved,Recorded,Scheduled,Posted,Skipped"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"N5:N{last}")
    ws.conditional_formatting.add(f"A5:P{last}", FormulaRule(formula=['$N5="Posted"'], fill=PatternFill("solid", fgColor="E3F1EE")))
    ws.conditional_formatting.add(f"A5:P{last}", FormulaRule(formula=['$D5="Month 2: Product launch"'], font=Font(color="7A3212")))
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:P{last}"
    s = wb.create_sheet("Summary")
    s["A1"] = "Counts by type and status (formulas)"
    s["A1"].font = Font(name=F, bold=True, size=12)
    types = sorted({r["type"] for r in rows})
    s["A3"], s["B3"], s["C3"], s["D3"] = "Type", "Planned", "Posted", "Remaining"
    for j in range(4):
        s.cell(row=3, column=1 + j).font = Font(name=F, bold=True)
    for k, t in enumerate(types):
        i = 4 + k
        s[f"A{i}"] = t
        s[f"B{i}"] = f'=COUNTIF(Calendar!$F$5:$F${last},A{i})'
        s[f"C{i}"] = f'=COUNTIFS(Calendar!$F$5:$F${last},A{i},Calendar!$N$5:$N${last},"Posted")'
        s[f"D{i}"] = f"=B{i}-C{i}"
    t = 4 + len(types)
    s[f"A{t}"], s[f"B{t}"], s[f"C{t}"], s[f"D{t}"] = "Total", f"=SUM(B4:B{t-1})", f"=SUM(C4:C{t-1})", f"=SUM(D4:D{t-1})"
    s[f"A{t}"].font = Font(name=F, bold=True)
    s.column_dimensions["A"].width = 32
    wb.save(path)


# ------------------------------------------------------------------ graphics

BASE = f"""{font_face_css()}*{{box-sizing:border-box}}html,body{{margin:0;overflow:hidden;font-family:Inter,sans-serif}}"""


def slide_html(text, n, total, ring, first=False, last=False):
    dark = first or last
    bg, fg = (C["ink"], C["paper"]) if dark else (C["paper"], C["ink"])
    size = 84 if first else (64 if len(text) < 60 else 54)
    ring_html = f'<div style="position:absolute;right:-240px;bottom:-240px;width:640px;height:640px;border-radius:50%;border:54px solid {ring}"></div>' if dark else \
        f'<div style="position:absolute;right:-120px;top:-120px;width:300px;height:300px;border-radius:50%;border:30px solid {ring};opacity:.9"></div>'
    extra = '<div style="font-size:30px;font-weight:600;color:%s;margin-top:34px">Save this · swipe →</div>' % C["butter"] if first else ""
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{BASE}
html,body{{width:1080px;height:1350px;background:{bg};color:{fg}}}
.t{{position:absolute;left:96px;right:120px;top:50%;transform:translateY(-55%);font-family:Fraunces,serif;font-weight:700;font-size:{size}px;line-height:1.12;letter-spacing:-.01em}}
.n{{position:absolute;right:80px;bottom:70px;font-size:26px;font-weight:700;color:{C['slate'] if not dark else '#D9D2C5'}}}
.logo{{position:absolute;left:96px;bottom:62px;height:48px}}
</style></head><body>{ring_html}<div class="t">{html.escape(text)}{extra}</div>
<img class="logo" src="{LOGO_DARK if dark else LOGO}"><div class="n">{n}/{total}</div></body></html>"""


def promo_html(text, title, product, ring):
    body = html.escape(card_text(text))
    shadow = "box-shadow:0 20px 50px rgba(0,0,0,.5);border-radius:4px;position:absolute"
    if product == "bundle":
        img = "".join(f'<img src="{COVERS[k]}" style="{shadow};right:{r}px;bottom:{b}px;width:250px;transform:rotate({d}deg)">'
                      for k, r, b, d in (("playbook", 250, 230, -6), ("desk", 170, 180, 0), ("prompts", 90, 130, 6)))
    elif COVERS.get(product):
        img = f'<img src="{COVERS[product]}" style="{shadow};right:90px;bottom:220px;width:330px;transform:rotate(4deg)">'
    else:
        img = ""
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{BASE}
html,body{{width:1080px;height:1350px;background:{C['ink']};color:{C['paper']}}}
.k{{position:absolute;left:90px;top:100px;font-weight:700;letter-spacing:.14em;text-transform:uppercase;font-size:26px;color:{C['butter']}}}
.t{{position:absolute;left:90px;right:90px;top:160px;font-family:Fraunces,serif;font-weight:700;font-size:78px;line-height:1.08}}
.b{{position:absolute;left:90px;width:{460 if product == "bundle" else 540}px;top:470px;font-size:36px;line-height:1.4;color:#EDE8DF}}
.logo{{position:absolute;left:90px;bottom:80px;height:54px}}
</style></head><body>
<div style="position:absolute;left:-260px;bottom:-300px;width:700px;height:700px;border-radius:50%;border:56px solid {ring}"></div>
<div class="k">Band of One</div><div class="t">{html.escape(title)}</div><div class="b">{body}</div>{img}
<img class="logo" src="{LOGO_DARK}"></body></html>"""


def card_text(text, limit=230):
    """First substantial paragraph, cut at a sentence (or word) boundary."""
    paras = [p.strip() for p in text.strip().split("\n\n") if p.strip()]
    para = next((p for p in paras if len(p) > 40 and not p.startswith(("•", "→", "✓", "["))), paras[0])
    if para.endswith(":"):  # a lead-in to a list: show the list items instead of a dangling colon
        i = paras.index(para)
        items = [ln.strip("•→✓ -").strip() for ln in (paras[i + 1] if i + 1 < len(paras) else "").splitlines() if ln.strip()]
        para = para + " " + " · ".join(items)
    para = " ".join(para.split())
    if len(para) <= limit:
        return para
    cut = para[:limit]
    end = max(cut.rfind(". "), cut.rfind(": "))
    return cut[: end + 1] if end > 60 else cut[: cut.rfind(" ")] + "…"


def cover_html(on_screen, concept, ring):
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{BASE}
html,body{{width:1080px;height:1920px;background:{C['ink']};color:{C['paper']}}}
.t{{position:absolute;left:90px;right:90px;top:520px;font-family:Fraunces,serif;font-weight:700;font-size:{96 if len(on_screen) < 45 else 76}px;line-height:1.06}}
.s{{position:absolute;left:90px;right:90px;top:360px;font-weight:700;letter-spacing:.12em;text-transform:uppercase;font-size:30px;color:{C['butter']}}}
.logo{{position:absolute;left:90px;bottom:160px;height:60px}}
</style></head><body>
<div style="position:absolute;right:-300px;bottom:-200px;width:900px;height:900px;border-radius:50%;border:70px solid {ring}"></div>
<div style="position:absolute;right:-120px;bottom:-40px;width:560px;height:560px;border-radius:50%;background:{C['teal']}"></div>
<div class="s">{html.escape(concept)}</div><div class="t">{html.escape(on_screen)}</div>
<img class="logo" src="{LOGO_DARK}"></body></html>"""


def graphics(base, videos, posts):
    g = base / "graphics"
    shutil.rmtree(g, ignore_errors=True)
    jobs = []
    for p in posts["educational"]:
        if not p.get("slides"):
            continue
        total = len(p["slides"])
        for i, s in enumerate(p["slides"], 1):
            h = slide_html(s, i, total, RING[post_tag(p)], first=i == 1, last=i == total)
            jobs.append({"html": write_html(f"sl_{p['id']}_{i}.html", h), "out": g / "carousels" / f"{p['id']}_{i:02d}.png",
                         "type": "png", "width": 1080, "height": 1350, "scale": 1})
    for p in posts["promotional"]:
        h = promo_html(p["text"], p["title"], p.get("product") or ("starter" if "/newsletter/" in p.get("link", "") else None), RING[post_tag(p)])
        jobs.append({"html": write_html(f"pr_{p['id']}.html", h), "out": g / "promo_cards" / f"{p['id']}.png",
                     "type": "png", "width": 1080, "height": 1350, "scale": 1})
    for v in videos:
        h = cover_html(v["on_screen"], v["concept"], RING[v["tags"]])
        jobs.append({"html": write_html(f"vc_{v['id']}.html", h), "out": g / "video_covers" / f"{v['id']}.png",
                     "type": "png", "width": 1080, "height": 1920, "scale": 1})
    run(jobs)
    return len(jobs)


TAGS = {}


def main():
    global TAGS
    c_videos = load(CB / "source/videos.yaml")
    c_posts = load(CB / "source/posts.yaml")
    p_videos = load(PB / "source/launch_videos.yaml")
    p_posts = load(PB / "source/launch_posts.yaml")
    TAGS = c_posts["hashtags"]
    for name, vids, posts in (("content", c_videos, c_posts), ("launch", p_videos, p_posts)):
        assert len(vids) == 30 and len(posts["educational"]) == 15 and len(posts["promotional"]) == 10, name
        assert sorted(v["day"] for v in vids) == list(range(1, 31)), name
    (CB / "Month1_Video_Concepts_Scripts_Captions.md").write_text(video_doc(
        c_videos, "Month 1 (content): 30 short-form videos",
        "Concepts, hooks, scripts, on-screen text, captions and hashtags. Each video points to one article or the free starter kit. "
        "Works as talking-head or faceless (screen recording + voiceover). **Nothing has been posted.**"))
    (CB / "Month1_Educational_and_Promotional_Posts.md").write_text(posts_doc(
        c_posts, "Month 1 (content): 15 educational + 10 promotional posts",
        "LinkedIn text posts and Instagram carousels (slides rendered in `graphics/carousels`). Replace `[link]` with the article URL once the site is live. **Nothing has been posted.**"))
    (PB / "Month2_Launch_Video_Ideas_Scripts_Captions.md").write_text(video_doc(
        p_videos, "Month 2 (product launch): 30 short-form videos",
        "Runs after Month 1. Pre-launch (days 1–7), launch window (days 8–21, real end date), post-launch (days 22–30). "
        "About one video in three asks for a sale. Fill in `[LAUNCH END DATE]` before recording. **Nothing has been posted.**"))
    (PB / "Month2_Launch_Educational_and_Promotional_Posts.md").write_text(posts_doc(
        p_posts, "Month 2 (product launch): 15 educational + 10 promotional posts",
        "Educational posts teach with real product content; promotional posts state real prices, the true launch end date and the refund policy. **Nothing has been posted.**"))
    rows = calendar_rows("Month 1: Content", 0, c_videos, c_posts) + calendar_rows("Month 2: Product launch", 30, p_videos, p_posts)
    write_csv([r for r in rows if r["day"] <= 30], CB / "Month1_Calendar_scheduler_import.csv")
    write_csv([r for r in rows if r["day"] > 30], PB / "Month2_Launch_Calendar_scheduler_import.csv")
    calendar_xlsx(rows, CB / "BandOfOne_60-Day_Social_Calendar.xlsx")
    shutil.copy(CB / "BandOfOne_60-Day_Social_Calendar.xlsx", PB / "BandOfOne_60-Day_Social_Calendar.xlsx")
    tag_md = ["# Hashtag sets", "", "Chosen for relevance to each topic. No hashtag volume data was available, and none is claimed. "
              "Use 3–5 per post, rotate between sets, and drop any that attract off-topic audiences.", ""]
    for k, v in TAGS.items():
        tag_md.append(f"- **{k}:** {' '.join(v)}")
    (CB / "Hashtag_Sets.md").write_text("\n".join(tag_md) + "\n")
    print("docs + calendars written;", len(rows), "calendar rows")
    if "--no-graphics" not in sys.argv:
        n = graphics(CB, c_videos, c_posts) + graphics(PB, p_videos, p_posts)
        print(n, "graphics rendered")


if __name__ == "__main__":
    main()
