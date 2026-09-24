"""Gumroad images: cover (1280x720), thumbnail (600x600) and preview images for each product.

Every image is made from the real product files (rendered PDF pages and real spreadsheet screenshots).
Run after the product builders:  python3 Build_Tools/build_store_images.py
"""
import shutil
import subprocess

import pypdfium2 as pdfium
from openpyxl import load_workbook
from PIL import Image, ImageChops

from brand import COLORS as C, ROOT, font_face_css
from render import TMP, run, write_html

DP = ROOT / "Digital_Products_Business"
OUT = DP / "07_Gumroad_Storefront" / "images"
WORK = TMP / "store"
LOGO_DARK = (DP / "02_Branding/logo/band-of-one-horizontal-on-dark.svg").as_uri()
LOGO = (DP / "02_Branding/logo/band-of-one-horizontal.svg").as_uri()

PDFS = {
    "playbook": DP / "03_Digital_Guide/BandOfOne_Solo-Admin-Playbook_v1.0.pdf",
    "deskguide": DP / "04_Premium_Template/BandOfOne_The-Client-Desk_User-Guide_v1.0.pdf",
    "prompts": DP / "05_AI_Prompt_Library/BandOfOne_Prompt-Library_v1.0.pdf",
    "starter": DP / "08_Marketing_Campaign/Lead_Magnet/BandOfOne_Client-Admin-Starter-Kit_FREE.pdf",
}


def page_png(key, idx, scale=2.0):
    out = WORK / f"{key}_p{idx + 1}.png"
    if not out.exists():
        doc = pdfium.PdfDocument(str(PDFS[key]))
        doc[idx].render(scale=scale).to_pil().convert("RGB").save(out)
    return out.as_uri()


def find_page(key, needle, start=1):
    doc = pdfium.PdfDocument(str(PDFS[key]))
    norm = lambda t: "".join(ch for ch in t.lower() if ch.isalnum())
    for i in range(start, len(doc)):
        text = norm(doc[i].get_textpage().get_text_range())
        if text.startswith("contents"):
            continue  # the table of contents lists every heading
        if norm(needle) in text[:len(norm(needle)) + 25]:
            return i
    raise ValueError(f"{needle!r} not found in {key}")


def search_screenshot():
    """Real screenshot of the Prompt Library's Search tab, via LibreOffice PDF export."""
    out = WORK / "search_tab.png"
    src = WORK / "pl_search.xlsx"
    wb = load_workbook(DP / "05_AI_Prompt_Library/BandOfOne_Prompt-Library_v1.0_Searchable.xlsx")
    for n in wb.sheetnames:
        if n != "Search":
            wb[n].sheet_state = "hidden"
    ws = wb["Search"]
    ws.column_dimensions["F"].width = 22
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:F20"
    wb.active = wb.sheetnames.index("Search")
    wb.save(src)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(WORK), str(src)],
                   check=True, capture_output=True)
    img = pdfium.PdfDocument(str(WORK / "pl_search.pdf"))[0].render(scale=2.4).to_pil().convert("RGB")
    bbox = ImageChops.difference(img, Image.new("RGB", img.size, (255, 255, 255))).getbbox()
    img.crop((max(bbox[0] - 30, 0), max(bbox[1] - 30, 0), min(bbox[2] + 30, img.width), min(bbox[3] + 30, img.height))).save(out)
    return out.as_uri()


BASE_CSS = f"""{font_face_css()}
*{{box-sizing:border-box}}
html,body{{margin:0;width:1280px;height:720px;overflow:hidden}}
body{{font-family:Inter,sans-serif;color:{C['ink']}}}
.dark{{background:{C['ink']};color:{C['paper']}}}
.light{{background:{C['paper']}}}
.kicker{{font-weight:700;letter-spacing:.16em;text-transform:uppercase;font-size:15px;color:{C['butter']}}}
.light .kicker{{color:{C['rust']}}}
h1{{font-family:Fraunces,serif;font-weight:700;font-size:62px;line-height:1.02;letter-spacing:-.02em;margin:14px 0 16px}}
h1 em{{font-style:italic;font-weight:400;color:{C['butter']}}}
.light h1 em{{color:{C['teal']}}}
.sub{{font-size:21px;line-height:1.45;color:#EDE8DF;max-width:520px}}
.light .sub{{color:{C['slate']}}}
.chips{{display:flex;flex-wrap:wrap;gap:10px;margin-top:26px;max-width:560px}}
.chip{{border:1.5px solid rgba(251,248,243,.35);border-radius:999px;padding:7px 15px;font-size:15px;font-weight:600}}
.light .chip{{border-color:{C['stone']};background:#fff}}
.left{{position:absolute;left:64px;top:70px;width:560px}}
.logo{{position:absolute;left:64px;bottom:44px;height:40px}}
.ring{{position:absolute;border-radius:50%}}
.shot{{position:absolute;box-shadow:0 22px 50px rgba(0,0,0,.38),0 3px 10px rgba(0,0,0,.25);border-radius:4px;background:#fff}}
.cap{{position:absolute;left:0;right:0;top:0;height:112px;padding:30px 56px 0;display:flex;align-items:baseline;gap:18px}}
.cap b{{font-family:Fraunces,serif;font-size:36px;font-weight:700}}
.cap span{{font-size:19px;color:{C['slate']}}}
.dark .cap span{{color:#D9D2C5}}
.win{{position:absolute;background:#fff;border-radius:10px;overflow:hidden;box-shadow:0 22px 50px rgba(0,0,0,.30)}}
.win .bar{{height:30px;background:#EDEAE4;display:flex;align-items:center;gap:7px;padding-left:12px}}
.win .bar i{{width:11px;height:11px;border-radius:50%;background:#CFC8BC;display:block}}
.win img{{display:block;width:100%}}
"""


def page(body, cls="dark"):
    return f'<!doctype html><html><head><meta charset="utf-8"><style>{BASE_CSS}</style></head><body class="{cls}">{body}</body></html>'


def rings(ring, disc):
    return (f'<div class="ring" style="right:-260px;top:250px;width:760px;height:760px;border:56px solid {ring}"></div>'
            f'<div class="ring" style="right:-120px;top:390px;width:480px;height:480px;background:{disc}"></div>')


def cover_img(uri, left, top, h, rot=0, z=2):
    return f'<img class="shot" src="{uri}" style="left:{left}px;top:{top}px;height:{h}px;transform:rotate({rot}deg);z-index:{z}">'


def hero(kicker, title, sub, chips, art, ring, disc):
    chip_html = "".join(f'<span class="chip">{c}</span>' for c in chips)
    return page(f'{rings(ring, disc)}<div class="left"><div class="kicker">{kicker}</div><h1>{title}</h1>'
                f'<div class="sub">{sub}</div><div class="chips">{chip_html}</div></div>'
                f'<img class="logo" src="{LOGO_DARK}" alt="Band of One">{art}')


def thumb(cover_uri, ring):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{BASE_CSS} html,body{{width:600px;height:600px}}</style></head>'
            f'<body class="dark"><div class="ring" style="right:-170px;bottom:-170px;width:460px;height:460px;border:40px solid {ring}"></div>'
            f'<img class="shot" src="{cover_uri}" style="left:152px;top:44px;height:512px"></body></html>')


def preview(title, sub, art, cls="light"):
    return page(f'<div class="cap"><b>{title}</b><span>{sub}</span></div>{art}', cls)


def main():
    shutil.rmtree(WORK, ignore_errors=True)
    WORK.mkdir(parents=True)
    OUT.mkdir(parents=True, exist_ok=True)
    jobs = []

    def add(name, html, w=1280, h=720):
        jobs.append({"html": write_html(f"store_{name}.html", html), "out": OUT / f"{name}.png",
                     "type": "png", "width": w, "height": h, "scale": 1.5})

    # ---------- A: Playbook
    pb_cover = page_png("playbook", 0)
    p_contents = page_png("playbook", 1)
    p_audit = page_png("playbook", find_page("playbook", "Run your admin audit"))
    p_paid = page_png("playbook", find_page("playbook", "Workflow 5: Getting paid"))
    p_prop = page_png("playbook", find_page("playbook", "Workflow 2: Proposals"))
    p_week = page_png("playbook", find_page("playbook", "Appendix C: Weekly review worksheet"))
    p_life = page_png("playbook", find_page("playbook", "The client lifecycle system"))
    add("A_playbook_cover", hero("The Playbook · PDF + editable DOCX", "The Solo Admin<br><em>Playbook</em>",
                                 "Build a simple AI-assisted admin system for a one-person service business, one workflow at a time.",
                                 ["38 pages", "6 client workflows", "Worksheets", "17 prompts", "Editable DOCX"],
                                 cover_img(p_contents, 820, 120, 520, 5, 1) + cover_img(pb_cover, 700, 92, 540, -4, 2),
                                 C["orange"], C["teal"]))
    add("A_playbook_thumbnail", thumb(pb_cover, C["orange"]), 600, 600)
    add("A_preview_1_inside", preview("Inside the playbook", "The client lifecycle, from first inquiry to referral",
                                      cover_img(p_contents, 120, 132, 560, -2) + cover_img(p_life, 660, 132, 560, 2)))
    add("A_preview_2_worksheets", preview("Worksheets you fill in", "Admin audit and weekly review",
                                          cover_img(p_audit, 120, 132, 560, -2) + cover_img(p_week, 660, 132, 560, 2)))
    add("A_preview_3_workflows", preview("Step-by-step workflows", "With templates, prompts and review checklists",
                                         cover_img(p_prop, 120, 132, 560, -2) + cover_img(p_paid, 660, 132, 560, 2)))

    # ---------- B: Client Desk
    dash = (DP / "04_Premium_Template/screenshot_dashboard.png").as_uri()
    im = Image.open(DP / "04_Premium_Template/screenshot_clients.png")
    im.crop((0, 0, int(im.width * 0.5), min(im.height, int(im.width * 0.5 * 0.34)))).save(WORK / "clients_crop.png")
    clients = (WORK / "clients_crop.png").as_uri()
    startshot = (DP / "04_Premium_Template/screenshot_start.png").as_uri()
    g_cover = page_png("deskguide", 0)
    g_tabs = page_png("deskguide", find_page("deskguide", "Quick start"))
    win = lambda uri, l, t, w: f'<div class="win" style="left:{l}px;top:{t}px;width:{w}px"><div class="bar"><i></i><i></i><i></i></div><img src="{uri}"></div>'
    add("B_clientdesk_cover", hero("The Workbook · Excel + Google Sheets", "The Client<br><em>Desk</em>",
                                   "Clients, projects, invoices and payment reminders in one spreadsheet, with a dashboard that shows what needs attention.",
                                   ["10 tabs", "Dashboard", "Reminder stages A–E", "Sample data", "13-page guide"],
                                   win(dash, 640, 110, 600) + cover_img(g_cover, 1010, 380, 300, 4, 3), C["butter"], C["teal"]))
    add("B_clientdesk_thumbnail", thumb(g_cover, C["butter"]), 600, 600)
    add("B_preview_1_dashboard", preview("The Dashboard", "Follow-ups due, overdue invoices, pipeline and this week's numbers",
                                         win(dash, 150, 128, 980)))
    add("B_preview_2_pipeline", preview("Clients & Pipeline", "Stage, value, next action and date for every client",
                                        win(clients, 90, 128, 1100)))
    add("B_preview_3_guide", preview("Set up in about 10 minutes", "Start Here tab plus a 13-page user guide",
                                     win(startshot, 90, 140, 640) + cover_img(g_tabs, 790, 130, 560, 2)))

    # ---------- C: Prompt Library
    pl_cover = page_png("prompts", 0)
    pl_glance = page_png("prompts", find_page("prompts", "The 100 prompts at a glance"))
    pl_set = page_png("prompts", find_page("prompts", "Setup & Foundations"))
    pl_pay = page_png("prompts", find_page("prompts", "Invoicing & Getting Paid"))
    pl_pay2 = page_png("prompts", find_page("prompts", "Invoicing & Getting Paid") + 1)
    search = search_screenshot()
    add("C_prompts_cover", hero("The Prompt Library · PDF + spreadsheet", "The Band of One<br><em>Prompt Library</em>",
                                "100 detailed prompts for client work and admin, from first inquiry to final invoice.",
                                ["100 prompts", "11 categories", "Inputs + output format", "Search tab", "CSV for Notion"],
                                cover_img(pl_set, 830, 120, 520, 5, 1) + cover_img(pl_cover, 700, 92, 540, -4, 2),
                                C["rust"], C["teal"]))
    add("C_prompts_thumbnail", thumb(pl_cover, C["rust"]), 600, 600)
    add("C_preview_1_cards", preview("Every prompt, fully specified", "Objective · when to use · inputs · prompt · output · tip",
                                     cover_img(pl_pay, 120, 132, 560, -2) + cover_img(pl_pay2, 660, 132, 560, 2)))
    add("C_preview_2_search", preview("Search by keyword", "The spreadsheet version finds matching prompts as you type",
                                      win(search, 90, 150, 1100)))
    add("C_preview_3_index", preview("11 categories across the client lifecycle", "Setup, leads, proposals, onboarding, delivery, payment, offboarding and more",
                                     cover_img(pl_glance, 120, 132, 560, -2) + cover_img(pl_set, 660, 132, 560, 2)))

    # ---------- Bundle
    add("D_bundle_cover", hero("The Complete Bundle · all three products", "The Band of One<br><em>Complete Bundle</em>",
                               "The playbook, the workbook and the prompt library: one system for the admin side of a business of one.",
                               ["Playbook", "The Client Desk", "Prompt Library", "One download"],
                               cover_img(pb_cover, 650, 58, 390, -3, 1) + cover_img(g_cover, 800, 160, 390, 0, 2)
                               + cover_img(pl_cover, 950, 262, 390, 3, 3), C["teal"], C["rust"]))
    bundle_thumb = page(f'<div class="ring" style="right:-170px;bottom:-170px;width:460px;height:460px;border:40px solid {C["teal"]}"></div>'
                        + cover_img(pb_cover, 48, 36, 330, -3, 1) + cover_img(g_cover, 150, 132, 330, 0, 2) + cover_img(pl_cover, 252, 228, 330, 3, 3)
                        ).replace("width:1280px;height:720px", "width:600px;height:600px")
    add("D_bundle_thumbnail", bundle_thumb, 600, 600)
    add("D_preview_1_whats_inside", preview("What's in the bundle", "Three products that work together or on their own",
                                            cover_img(pb_cover, 110, 150, 470, -3) + cover_img(g_cover, 425, 140, 470, 0)
                                            + cover_img(pl_cover, 740, 150, 470, 3)))

    # ---------- Free starter kit (lead magnet)
    sk_cover = page_png("starter", 0)
    sk_check = page_png("starter", find_page("starter", "Your 30-minute weekly review"))
    sk_card = page_png("starter", find_page("starter", "The 10 prompts"))
    add("E_starter_cover", hero("Free · PDF", "The Client Admin<br><em>Starter Kit</em>",
                                "10 AI prompts for the admin you do every week, plus a 30-minute weekly review checklist.",
                                ["10 prompts", "Weekly checklist", "Free"],
                                cover_img(sk_check, 830, 120, 520, 5, 1) + cover_img(sk_cover, 700, 92, 540, -4, 2),
                                C["teal"], C["rust"]))
    add("E_starter_thumbnail", thumb(sk_cover, C["teal"]), 600, 600)
    add("E_preview_1_inside", preview("Inside the starter kit", "Ready-to-use prompts and a checklist you can print",
                                      cover_img(sk_card, 120, 132, 560, -2) + cover_img(sk_check, 660, 132, 560, 2)))
    run(jobs)
    # optimize PNGs (palette where lossless enough is not guaranteed, so keep RGB but optimize)
    for j in jobs:
        im = Image.open(j["out"]).convert("RGB")
        im.save(j["out"], optimize=True)
    print(f"{len(jobs)} images -> {OUT}")


if __name__ == "__main__":
    main()
