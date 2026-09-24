"""Build the free 'Invoice Tracker Starter' sheet used in the article
'How to Track Invoices and Reminders in Google Sheets or Excel' and test its formulas.

Run: python3 Build_Tools/build_invoice_starter.py          (build + test)
The formulas here are printed verbatim in the article; the test proves they work.
"""
import datetime as dt
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from brand import COLORS as C, ROOT

OUT = ROOT / "Content_Automation_Business" / "03_Website" / "src" / "downloads" / "BandOfOne_Invoice-Tracker-Starter.xlsx"
RECALC = next(Path("/root/.claude/skills").rglob("xlsx/scripts/recalc.py"), None)
ROWS = range(4, 54)

FORMULAS = {
    "F": '=IF(D{r}="","",D{r}+E{r})',
    "H": '=IF(D{r}="","",IF(G{r}<>"","Paid",IF($B$1<F{r},"Not due",IF($B$1=F{r},"Due today","Overdue"))))',
    "I": '=IF(OR(D{r}="",G{r}<>""),"",MAX(0,$B$1-F{r}))',
    "J": '=IF(OR(D{r}="",G{r}<>""),"",IF($B$1>=F{r}+30,"E",IF($B$1>=F{r}+14,"D",IF($B$1>=F{r}+7,"C",IF($B$1>=F{r},"B",IF($B$1>=F{r}-3,"A","—"))))))',
    "L": '=IF(OR(J{r}="",J{r}="—"),"",IF(K{r}="","Send "&J{r},IF(CODE(J{r})>CODE(K{r}),"Send "&J{r},"Up to date")))',
}


def build(out: Path, today=None, rows=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    F = "Arial"
    ws["A1"], ws["B1"] = "Today", today if today else "=TODAY()"
    ws["A1"].font = Font(name=F, bold=True)
    ws["B1"].number_format = "yyyy-mm-dd"
    ws["D1"] = "Yellow = type here · grey = calculated. Stages: A = 3 days before due · B = due date · C = +7 · D = +14 · E = +30 days."
    ws["D1"].font = Font(name=F, italic=True, size=9, color=C["slate"][1:])
    heads = ["Invoice #", "Client", "Amount", "Issue date", "Terms (days)", "Due date", "Paid date", "Status",
             "Days overdue", "Reminder stage due", "Last reminder sent", "Action"]
    widths = [11, 22, 11, 12, 11, 12, 12, 11, 12, 12, 14, 14]
    for j, (h, w) in enumerate(zip(heads, widths)):
        c = ws.cell(row=3, column=1 + j, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C["deep_teal"][1:])
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[chr(65 + j)].width = w
    yellow, grey = PatternFill("solid", fgColor="FFF8E1"), PatternFill("solid", fgColor="F2F2F2")
    for r in ROWS:
        for col in "ABCDEGK":
            ws[f"{col}{r}"].fill = yellow
        for col, f in FORMULAS.items():
            ws[f"{col}{r}"] = f.format(r=r)
            ws[f"{col}{r}"].fill = grey
        for col in "DFG":
            ws[f"{col}{r}"].number_format = "yyyy-mm-dd"
        ws[f"C{r}"].number_format = "#,##0.00"
        for col in "ABCDEFGHIJKL":
            ws[f"{col}{r}"].font = Font(name=F)
    for k, row in enumerate(rows or []):
        for col, v in zip("ABCDEGK", row):
            if v is not None:
                ws[f"{col}{4 + k}"] = v
    dv = DataValidation(type="list", formula1='"A,B,C,D,E"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("K4:K53")
    red = PatternFill("solid", fgColor="F8D7CD")
    ws.conditional_formatting.add("H4:H53", FormulaRule(formula=['$H4="Overdue"'], fill=red))
    ws.conditional_formatting.add("L4:L53", FormulaRule(formula=['LEFT($L4,4)="Send"'], fill=PatternFill("solid", fgColor="FCE7A8")))
    ws.freeze_panes = "A4"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def recalc(path):
    res = json.loads(subprocess.run([sys.executable, str(RECALC), str(path), "60"], capture_output=True, text=True).stdout)
    assert res.get("status") == "success", res
    return res


def test():
    T = dt.date(2026, 10, 15)
    d = lambda n: T + dt.timedelta(days=n)
    # (A..G,K inputs) and expected (F due, H status, I overdue, J stage, L action)
    cases = [
        (("1", "Paid on time", 500, d(-20), 14, d(-8), None), (d(-6), "Paid", None, None, None)),
        (("2", "Not due yet", 800, d(-2), 14, None, None), (d(12), "Not due", 0, "—", None)),
        (("3", "Heads-up window", 800, d(-11), 14, None, None), (d(3), "Not due", 0, "A", "Send A")),
        (("4", "Due today", 300, d(-14), 14, None, "A"), (d(0), "Due today", 0, "B", "Send B")),
        (("5", "8 days late", 1200, d(-22), 14, None, "C"), (d(-8), "Overdue", 8, "C", "Up to date")),
        (("6", "15 days late", 180, d(-29), 14, None, None), (d(-15), "Overdue", 15, "D", "Send D")),
        (("7", "45 days late", 900, d(-75), 30, None, "D"), (d(-45), "Overdue", 45, "E", "Send E")),
        (("8", "Exactly +7", 400, d(-37), 30, None, "B"), (d(-7), "Overdue", 7, "C", "Send C")),
    ]
    tmp = Path(tempfile.mkdtemp()) / "t.xlsx"
    build(tmp, today=T, rows=[c[0] for c in cases])
    recalc(tmp)
    ws = load_workbook(tmp, data_only=True).active
    fails = 0
    for k, (_, exp) in enumerate(cases):
        r = 4 + k
        due = ws[f"F{r}"].value
        due = due.date() if hasattr(due, "date") else due
        got = (due, ws[f"H{r}"].value, ws[f"I{r}"].value, ws[f"J{r}"].value, ws[f"L{r}"].value)
        got = tuple(None if v == "" else v for v in got)
        ok = got == exp
        fails += not ok
        print("PASS" if ok else "FAIL", ws[f"B{r}"].value, got if not ok else "")
    empty = [ws[f"{c}20"].value for c in "FHIJL"]
    ok = all(v in (None, "") for v in empty)
    fails += not ok
    print("PASS" if ok else "FAIL", "empty rows stay blank")
    print(f"{len(cases) + 1 - fails}/{len(cases) + 1} passed")
    return fails


if __name__ == "__main__":
    failures = test()
    build(OUT, rows=[("INV-001", "Example client (delete me)", 500, "=B1-10", 14, None, None)])
    recalc(OUT)
    print(OUT)
    sys.exit(1 if failures else 0)
