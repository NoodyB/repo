"""Affiliate opportunities workbook: verified program terms, application tracker, link placeholder map,
disclosure language. NO affiliate links are included: none have been issued.

Run: python3 Build_Tools/build_affiliate_workbook.py
"""
import datetime as dt

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from brand import COLORS as C, ROOT

OUT = ROOT / "Content_Automation_Business" / "06_Affiliate_Marketing" / "BandOfOne_Affiliate_Opportunities.xlsx"
V = "2026-09-24"
NS = "Not stated on page: check terms at signup"
F = "Arial"
thin = Side(style="thin", color="D9D4CA")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor=C["deep_teal"][1:])
INPUT = PatternFill("solid", fgColor="FFF8E1")
CALC = PatternFill("solid", fgColor="F2F2F2")
NOT_ISSUED = "[NOT ISSUED: do not use]"

# company, program, official page, how to apply, commission, cookie, payout threshold/method, geo, approval, topics, program disclosure notes, verification
PROGRAMS = [
    ("MailerLite", "MailerLite Affiliate Program", "https://www.mailerlite.com/affiliate", "mailerlite.trackdesk.com/sign-up",
     "30% recurring (lifetime of referral)", "45 days", "$100 minimum open balance from at least 2 unique sales; settled Fridays; 30-day on-hold period",
     NS, "Application reviewed (up to 10 business days)", "Newsletter tools; mailerlite-vs-kit-vs-beehiiv; minimum-tool-stack", "Program terms page: mailerlite.com/legal/affiliate-program-terms", "Verified official page"),
    ("Kit", "Kit Affiliate Program", "https://kit.com/affiliate", "dash.partnerstack.com/application?company=kit&group=kitaffiliates",
     "50% for first 12 months; 10–20% recurring after 12 months with status tiers", "Not stated", "$5 threshold; paid around the 13th monthly (Stripe/PayPal/direct deposit); 31-day hold",
     NS, "Applications reviewed individually; no audience-size requirement stated", "Newsletter tools; creator email", "Via PartnerStack", "Verified official page"),
    ("beehiiv", "beehiiv Partner Program", "https://www.beehiiv.com/partners/affiliate", "app.beehiiv.com/partner_program",
     "Up to 60% of referral's payments for 12 months", "60 days", "Paid on the 15th for the prior month (threshold not captured)",
     NS, NS, "Newsletter platforms comparison", "Managed via Dub.co", "Verified official page"),
    ("Notion", "Notion Affiliate Program", "https://www.notion.com/affiliates", "Via PartnerStack (link on official page)",
     "Up to $50 per activated signup + 20% of year-one revenue", "180 days", NS,
     NS, "Application (audience types listed on page)", "Client tracking; spreadsheet-vs-notion-vs-crm; SOP storage", "Via PartnerStack", "Verified official page"),
    ("Make", "Make Affiliate Program", "https://www.make.com/en/affiliate", "make.com/user/affiliate",
     "35% for 12 months", "30 days", "$100 minimum; paid via Wise",
     NS, "Open to anyone with a Make account (per page)", "Automation; make-vs-zapier-vs-n8n; what-to-automate-first", "", "Verified official page"),
    ("n8n", "n8n Affiliate Program", "https://n8n.io/affiliates", "dash.partnerstack.com/application?company=n8n&group=affiliates2025",
     "30% of n8n Cloud referrals for the first year", "Not stated", "Monthly payouts ≥ EUR 100 via PayPal",
     NS, "Application via PartnerStack", "Automation comparison", "Cloud plans only (self-hosted is free)", "Verified official page"),
    ("Tally", "Tally Referral Program", "https://tally.so/help/referral-program", "Tally dashboard → Rewards (needs a Tally account)",
     "20% of subscription payments, capped at $150 per referred user", "Not stated", "PayPal or Venmo (threshold not captured)",
     NS, "No paid plan needed; self-referrals not allowed", "Intake forms; client-intake-form-tools; onboarding", "Referral program (not a classic affiliate network)", "Verified official page"),
    ("Dubsado", "Dubsado Affiliate Program", "https://www.dubsado.com/affiliate-program", "app.dubsado.com/user/signup",
     "$35 per referral that stays on a paid plan for 60 days", "Not stated", NS,
     NS, "Open to anyone, including free-trial users (per page)", "Client management (CRM) comparisons", "", "Verified official page"),
    ("Bonsai", "Bonsai Affiliate Program", "https://www.hellobonsai.com/affiliates", "hellobonsai.firstpromoter.com",
     "200% of first monthly payment (monthly plans); 30% (yearly plans)", "60 days", "Monthly payouts (threshold not captured)",
     NS, NS, "Invoicing/CRM for freelancers", "Via FirstPromoter", "Verified official page"),
    ("HoneyBook", "HoneyBook Referral Program", "https://www.honeybook.com/partnerships/referrals", "Via HoneyBook account",
     "$100 per paid subscriber + milestone bonuses; described as a limited-time offer", "Not stated", "1099-NEC issued if ≥ $600/year (US)",
     "US-focused (verify)", "Existing HoneyBook user referral", "Client management comparisons", "'Limited time offer, terms apply': re-check before relying on it", "Verified official page"),
    ("Buffer", "Buffer Partner/Affiliate Program", "https://buffer.com/partners", "Via official page",
     "25% on new customers for 12 months (only when the referral pays)", "Not stated", NS,
     NS, NS, "Social scheduling (marketing articles, future)", "", "Verified official page"),
    ("Hostinger", "Hostinger Affiliate Program", "https://www.hostinger.com/affiliates", "affiliates.hostinger.com",
     "At least 40% per eligible sale; rises with volume", "Not stated", NS,
     NS, NS, "Website hosting (only if you write hosting content)", "", "Verified official page"),
    ("Descript", "Descript Affiliate Program", "https://www.descript.com/affiliate", "Via PartnerStack",
     "$25 flat per new qualifying subscription", "Not stated", "Paid via PayPal",
     NS, NS, "Video/podcast editing (social content workflow)", "Terms: descript.com/affiliate-terms", "Verified official page"),
    ("Amazon", "Amazon Associates (US)", "https://affiliate-program.amazon.com/", "affiliate-program.amazon.com",
     "Fixed rate table by category, e.g. Physical Books 4.50%", "Per Amazon policy", "Per Amazon policy",
     "Per-marketplace programs", "Application; sales requirements apply (check operating agreement)", "Book recommendations (business books)", "Amazon requires specific disclosure wording: check the Operating Agreement", "Verified official page (rate table)"),
    ("Fathom", "Fathom Growth Partner Program", "https://fathom.ai/program/growth-partner", "Via PartnerStack",
     "Terms not captured", "Not captured", "Not captured", NS, "For agencies and consultants", "AI meeting notes (future)", "", "Partially verified: page exists, terms not captured"),
]

NOT_AVAILABLE = [
    ("Canva", "Affiliate program replaced by 'Canvassador'; closed to applications when checked", "canva.com/help/canva-affiliate-marketing-program"),
    ("Zapier", "No public affiliate page found (zapier.com/affiliates returned 404)", "zapier.com/affiliates"),
    ("Motion", "Affiliate page returned 404", "usemotion.com/affiliate-program"),
    ("Jotform", "Affiliate page returned 404", "jotform.com/affiliate-program/"),
    ("ClickUp", "Site blocked our check (403). A third-party summary claimed terms we could NOT verify: treat as unknown", "clickup.com"),
    ("ChatGPT / Claude / Gemini", "No consumer affiliate programs identified; not researched further", "-"),
    ("Google Forms", "No affiliate program (free Google product)", "-"),
]

# article slug, tool, official link used today, placeholder id
PLACEMENTS = [
    ("mailerlite-vs-kit-vs-beehiiv", "MailerLite", "https://www.mailerlite.com/pricing", "AFF_MAILERLITE"),
    ("mailerlite-vs-kit-vs-beehiiv", "Kit", "https://kit.com/pricing", "AFF_KIT"),
    ("mailerlite-vs-kit-vs-beehiiv", "beehiiv", "https://www.beehiiv.com/pricing", "AFF_BEEHIIV"),
    ("minimum-tool-stack", "MailerLite", "https://www.mailerlite.com/pricing", "AFF_MAILERLITE"),
    ("minimum-tool-stack", "Kit", "https://kit.com/pricing", "AFF_KIT"),
    ("minimum-tool-stack", "Tally", "https://tally.so/pricing", "AFF_TALLY"),
    ("minimum-tool-stack", "Notion", "https://www.notion.com/pricing", "AFF_NOTION"),
    ("minimum-tool-stack", "Make", "https://www.make.com/en/pricing", "AFF_MAKE"),
    ("make-vs-zapier-vs-n8n", "Make", "https://www.make.com/en/pricing", "AFF_MAKE"),
    ("make-vs-zapier-vs-n8n", "n8n", "https://n8n.io/pricing/", "AFF_N8N"),
    ("what-to-automate-first", "Make", "https://www.make.com/en/pricing", "AFF_MAKE"),
    ("client-intake-form-tools", "Tally", "https://tally.so/pricing", "AFF_TALLY"),
    ("spreadsheet-vs-notion-vs-crm", "Notion", "https://www.notion.com/pricing", "AFF_NOTION"),
    ("client-onboarding-checklist", "Tally", "https://tally.so/pricing", "AFF_TALLY"),
]

DISCLOSURES = [
    ("Site-wide page", "Affiliate disclosure page", "Already live on the website draft: /affiliate-disclosure/ (currently states that there are no affiliate links)."),
    ("Top of article (before first link)", "Article disclosure box",
     "Disclosure: this article contains affiliate links, marked \"(affiliate link)\". If you buy through one, we may earn a commission at no extra cost to you. It never changes our recommendations."),
    ("Next to each link", "Inline label", "Tool name (affiliate link)"),
    ("Tool comparison intro", "Relationship statement",
     "We're an affiliate for [TOOL A] and [TOOL B], which means we may earn a commission if you sign up through our links. We have no relationship with [TOOL C]. Our ratings are the same either way."),
    ("Email newsletter", "Email line", "This email contains an affiliate link to [TOOL]. If you sign up through it, I may earn a commission at no extra cost to you."),
    ("Social post (text)", "Post disclosure", "#ad: I earn a commission if you sign up through this link. I use/recommend it because [honest reason]."),
    ("Video / short-form", "Spoken + on-screen", "Say it out loud near the start (\"This video includes an affiliate link\") AND show it on screen; a caption alone may be missed."),
    ("Own products", "Own-product statement", "This is a product we make and sell."),
    ("Amazon links", "Amazon required wording", "As an Amazon Associate I earn from qualifying purchases. (Check the current Operating Agreement for the exact wording required.)"),
]


def style_header(ws, row, heads, widths):
    for j, (h, w) in enumerate(zip(heads, widths)):
        c = ws.cell(row=row, column=1 + j, value=h)
        c.font, c.fill, c.border = Font(name=F, bold=True, color="FFFFFF"), HEAD, B
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w


def body(ws, r0, r1, c1):
    for row in ws.iter_rows(min_row=r0, max_row=r1, max_col=c1):
        for c in row:
            c.font = Font(name=F, size=9 if c.column > 2 else 10, bold=c.column == 1)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = B


def main():
    wb = Workbook()
    rd = wb.active
    rd.title = "Read Me"
    rd.sheet_view.showGridLines = False
    rd.column_dimensions["A"].width = 3
    rd.column_dimensions["B"].width = 120
    lines = [
        ("Band of One: Affiliate opportunities & disclosure kit", Font(name=F, bold=True, size=16, color=C["deep_teal"][1:])),
        (f"Terms captured from each program's official page on {V}. Re-verify before applying and every 90 days (see 'Days since verified').", Font(name=F, italic=True, color=C["slate"][1:])),
        ("", None),
        ("RULES", Font(name=F, bold=True, size=12)),
        ("1. You have NOT been accepted into any program. Every status starts as 'Not applied'.", Font(name=F)),
        (f"2. Never put an affiliate link on the site, in emails or in posts until the program has issued it to you. Until then the 'Issued link' column says {NOT_ISSUED}.", Font(name=F)),
        ("3. When a link is issued: paste it into 'Issued link', set Status = Approved, then use the 'Link Placeholder Map' tab to see which articles can be updated.", Font(name=F)),
        ("4. Every article with an affiliate link gets `affiliate: true` in its front matter (this shows the disclosure box) and '(affiliate link)' next to each link. Use rel=\"sponsored\".", Font(name=F)),
        ("5. Commission never decides a recommendation. If a free or non-affiliate option is better, say so.", Font(name=F)),
        ("6. Don't use affiliate links in places a program forbids (some ban links in paid ads, emails or on coupon sites). Read each program's terms.", Font(name=F)),
        ("7. FTC guidance: disclose material connections clearly and conspicuously: https://www.ftc.gov/business-guidance/resources/ftcs-endorsement-guides-what-people-are-asking", Font(name=F)),
        ("", None),
        ("TABS", Font(name=F, bold=True, size=12)),
        ("Programs: verified terms + your application tracker (yellow = you fill in).", Font(name=F)),
        ("Link Placeholder Map: where an issued link would replace an official link, with a formula that says whether it's allowed yet.", Font(name=F)),
        ("Disclosure Language: copy-paste disclosures for pages, articles, email, social and video.", Font(name=F)),
        ("Not Available: programs that were closed, missing or unverifiable when checked.", Font(name=F)),
        ("", None),
        ("No commission amounts, conversion rates or earnings are projected anywhere in this workbook.", Font(name=F, italic=True, color=C["rust"][1:])),
    ]
    for k, (t, f) in enumerate(lines, start=2):
        rd[f"B{k}"] = t
        if f:
            rd[f"B{k}"].font = f
        rd[f"B{k}"].alignment = Alignment(wrap_text=True)

    # Programs
    p = wb.create_sheet("Programs")
    heads = ["Company", "Program", "Official program page", "How to apply", "Commission (as published)", "Cookie / attribution window",
             "Payout threshold & method", "Geographic restrictions", "Approval requirements", "Relevant content topics",
             "Program-specific notes / disclosure rules", "Verification", "Last verified", "Days since verified", "Re-verify?",
             "Status", "Date applied", "Issued link", "Notes"]
    widths = [12, 20, 30, 28, 30, 13, 30, 20, 26, 28, 28, 18, 12, 10, 10, 12, 12, 22, 24]
    style_header(p, 1, heads, widths)
    p.row_dimensions[1].height = 42
    for k, row in enumerate(PROGRAMS):
        r = 2 + k
        for j, v in enumerate(row):
            p.cell(row=r, column=1 + j, value=v)
        p.cell(row=r, column=13, value=dt.date(2026, 9, 24)).number_format = "yyyy-mm-dd"
        p.cell(row=r, column=14, value=f'=IF(M{r}="","",TODAY()-M{r})')
        p.cell(row=r, column=15, value=f'=IF(N{r}="","",IF(N{r}>90,"Re-verify","OK"))')
        p.cell(row=r, column=16, value="Not applied")
        p.cell(row=r, column=18, value=NOT_ISSUED)
    last = 1 + len(PROGRAMS)
    body(p, 2, last, len(heads))
    for r in range(2, last + 1):
        for col in "PQRS":
            p[f"{col}{r}"].fill = INPUT
        for col in "NO":
            p[f"{col}{r}"].fill = CALC
    dv = DataValidation(type="list", formula1='"Not applied,Applied,Approved,Rejected,Paused"', allow_blank=False)
    p.add_data_validation(dv)
    dv.add(f"P2:P{last}")
    p.conditional_formatting.add(f"O2:O{last}", FormulaRule(formula=['$O2="Re-verify"'], fill=PatternFill("solid", fgColor="F8D7CD")))
    p.conditional_formatting.add(f"P2:P{last}", FormulaRule(formula=['$P2="Approved"'], fill=PatternFill("solid", fgColor="D8EFE9")))
    p.freeze_panes = "C2"
    p.auto_filter.ref = f"A1:S{last}"

    # Placeholder map
    m = wb.create_sheet("Link Placeholder Map")
    heads = ["Article slug", "Tool", "Official link used today", "Placeholder ID", "Program status", "Issued link", "Action"]
    style_header(m, 1, heads, [30, 12, 36, 18, 14, 26, 44])
    for k, (slug, tool, url, pid) in enumerate(PLACEMENTS):
        r = 2 + k
        m[f"A{r}"], m[f"B{r}"], m[f"C{r}"], m[f"D{r}"] = slug, tool, url, pid
        m[f"E{r}"] = f'=IFERROR(INDEX(Programs!$P$2:$P${last},MATCH(B{r},Programs!$A$2:$A${last},0)),"No program")'
        m[f"F{r}"] = f'=IFERROR(INDEX(Programs!$R$2:$R${last},MATCH(B{r},Programs!$A$2:$A${last},0)),"")'
        m[f"G{r}"] = (f'=IF(AND(E{r}="Approved",F{r}<>"",F{r}<>"{NOT_ISSUED}"),'
                      f'"Allowed: swap link, label (affiliate link), set affiliate: true","Keep the official link")')
    lm = 1 + len(PLACEMENTS)
    body(m, 2, lm, 7)
    for r in range(2, lm + 1):
        for col in "EFG":
            m[f"{col}{r}"].fill = CALC
    m.freeze_panes = "A2"
    m.auto_filter.ref = f"A1:G{lm}"

    d = wb.create_sheet("Disclosure Language")
    style_header(d, 1, ["Where", "Type", "Copy-paste text"], [26, 24, 110])
    for k, row in enumerate(DISCLOSURES):
        for j, v in enumerate(row):
            d.cell(row=2 + k, column=1 + j, value=v)
    body(d, 2, 1 + len(DISCLOSURES), 3)

    n = wb.create_sheet("Not Available")
    style_header(n, 1, ["Company", "Finding (checked " + V + ")", "Page checked"], [22, 90, 40])
    for k, row in enumerate(NOT_AVAILABLE):
        for j, v in enumerate(row):
            n.cell(row=2 + k, column=1 + j, value=v)
    body(n, 2, 1 + len(NOT_AVAILABLE), 3)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = C["teal"][1:]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
